#!/usr/bin/env python3
"""Inserta diapositivas de k-NN (UD3 UAX → recuperación densa AgroPS)."""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "Defensa_TFM_AgroPS_15min.pptx"
ALT = ROOT / "Defensa_TFM_AgroPS_15min_knn.pptx"

BG = RGBColor(0xF7, 0xF5, 0xF0)
INK = RGBColor(0x1A, 0x2E, 0x1A)
MUTED = RGBColor(0x4A, 0x5C, 0x4A)
ACCENT = RGBColor(0x2F, 0x6B, 0x3A)
ACCENT2 = RGBColor(0xC4, 0x7A, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xE8, 0xEF, 0xE6)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARKER = "k-NN: de la clasificación (UD3) a la recuperación RAG"


def _fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    _fill(shape, color)
    return shape


def _round(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    _fill(shape, color)
    return shape


def _set_run(run, text, size=18, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _textbox(slide, left, top, width, height, specs, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[valign],
        )
    except Exception:
        pass
    first = True
    for spec in specs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = spec.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(spec.get("space_after", 4))
        run = p.add_run()
        _set_run(
            run,
            spec["text"],
            size=spec.get("size", 14),
            bold=spec.get("bold", False),
            color=spec.get("color", INK),
            font=spec.get("font", "Calibri"),
        )
    return box


def _badge(slide, left, top, width, height, text, bg=ACCENT, fg=WHITE, size=12):
    shape = _round(slide, left, top, width, height, bg)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    _set_run(run, text, size=size, bold=True, color=fg)
    try:
        tf._txBody.bodyPr.set("anchor", "ctr")
    except Exception:
        pass
    return shape


def _title_bar(slide, title, subtitle=None):
    _rect(slide, 0, 0, SLIDE_W, Inches(0.08), ACCENT)
    _textbox(
        slide,
        Inches(0.5),
        Inches(0.22),
        Inches(12.3),
        Inches(0.5),
        [{"text": title, "size": 24, "bold": True, "color": INK}],
    )
    if subtitle:
        _textbox(
            slide,
            Inches(0.5),
            Inches(0.72),
            Inches(12.3),
            Inches(0.32),
            [{"text": subtitle, "size": 13, "color": MUTED}],
        )


def _footer(slide, note="k-NN"):
    _textbox(
        slide,
        Inches(0.4),
        Inches(7.05),
        Inches(10),
        Inches(0.32),
        [{"text": f"AgroPS · Defensa TFM · {note}", "size": 10, "color": MUTED}],
    )


def _move_slide(prs: Presentation, old_index: int, new_index: int) -> None:
    sld_id_lst = prs.slides._sldIdLst  # noqa: SLF001
    slides = list(sld_id_lst)
    el = slides[old_index]
    sld_id_lst.remove(el)
    sld_id_lst.insert(new_index, el)


def _has_knn(prs: Presentation) -> bool:
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame and MARKER in sh.text_frame.text:
                return True
    return False


def _find_insert_after(prs: Presentation) -> int:
    for i, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            if sh.has_text_frame and "Técnicas de IA aplicadas" in sh.text_frame.text:
                return i + 1
    return 7


def _enrich_embeddings_card(prs: Presentation) -> None:
    old = "nomic-embed-text → pgvector (kNN coseno)."
    new = "k-NN lazy: pregunta vs chunks. Distancia coseno en pgvector/HNSW (k=10)."
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for run in p.runs:
                    if run.text.strip() == old:
                        run.text = new


def add_knn_slides(prs: Presentation, insert_at: int) -> None:
    blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]

    # ── Slide A: k-NN UD3 vs AgroPS ─────────────────────────────
    s = prs.slides.add_slide(blank)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, BG)
    _title_bar(
        s,
        MARKER,
        "Mismo learner instance-based; distinto objetivo (recuperar evidencia, no votar una clase)",
    )

    _round(s, Inches(0.4), Inches(1.18), Inches(6.15), Inches(4.55), CARD)
    _rect(s, Inches(0.4), Inches(1.18), Inches(0.12), Inches(4.55), ACCENT2)
    _badge(s, Inches(0.75), Inches(1.38), Inches(3.6), Inches(0.38), "UD3 · Clasificación", ACCENT2)
    _textbox(
        s,
        Inches(0.75),
        Inches(1.9),
        Inches(5.55),
        Inches(3.6),
        [
            {
                "text": "k-NN es supervisado, no paramétrico y lazy: no estima un modelo; almacena los ejemplos y decide en consulta.",
                "size": 13,
                "color": INK,
                "space_after": 10,
            },
            {"text": "1. Elegir k (sesgo–varianza).", "size": 14, "color": INK, "space_after": 4},
            {"text": "2. Medir distancia a todos los casos (euclídea en el curso).", "size": 14, "color": INK, "space_after": 4},
            {"text": "3. Tomar los k vecinos más cercanos.", "size": 14, "color": INK, "space_after": 4},
            {"text": "4. Predecir por voto mayoritario de la etiqueta.", "size": 14, "color": INK, "space_after": 10},
            {
                "text": "Ejemplo del curso: maligno/benigno. El output es una clase.",
                "size": 12,
                "color": MUTED,
            },
        ],
    )

    _round(s, Inches(6.75), Inches(1.18), Inches(6.15), Inches(4.55), CARD)
    _rect(s, Inches(6.75), Inches(1.18), Inches(0.12), Inches(4.55), ACCENT)
    _badge(s, Inches(7.1), Inches(1.38), Inches(3.9), Inches(0.38), "AgroPS · Recuperación densa", ACCENT)
    _textbox(
        s,
        Inches(7.1),
        Inches(1.9),
        Inches(5.55),
        Inches(3.6),
        [
            {
                "text": "Los “ejemplos” son embeddings de chunks (nomic-embed-text). “Entrenar” es persistirlos en pgvector.",
                "size": 13,
                "color": INK,
                "space_after": 10,
            },
            {"text": "1. k = retrieval_k = 10.", "size": 14, "color": INK, "space_after": 4},
            {"text": "2. Distancia coseno pregunta ↔ chunk.", "size": 14, "color": INK, "space_after": 4},
            {"text": "3. Devolver los k chunks más cercanos (no hay voto).", "size": 14, "color": INK, "space_after": 4},
            {"text": "4. Esos fragmentos se inyectan al prompt de Llama 3.2.", "size": 14, "color": INK, "space_after": 10},
            {
                "text": "La etiqueta no existe: el “vecino” es evidencia documental.",
                "size": 12,
                "color": MUTED,
            },
        ],
    )

    _round(s, Inches(0.4), Inches(5.88), Inches(12.5), Inches(0.95), SOFT)
    _textbox(
        s,
        Inches(0.65),
        Inches(6.05),
        Inches(12.1),
        Inches(0.7),
        [
            {
                "text": "Lazy learner: el coste está en la consulta (medir distancias). El curso recomienda indexación / vecinos aproximados → en AgroPS, HNSW con vector_cosine_ops.",
                "size": 14,
                "bold": True,
                "color": INK,
            }
        ],
    )
    _footer(s, "k-NN")

    # ── Slide B: métrica coseno ─────────────────────────────────
    s = prs.slides.add_slide(blank)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, BG)
    _title_bar(
        s,
        "Métrica de búsqueda: distancia coseno",
        "UD3: euclídea / Manhattan / coseno. Producción AgroPS: solo coseno (L1/L2 quedan para el laboratorio)",
    )

    metrics = [
        (
            "Coseno (producción)",
            "1 − (q · d) / (‖q‖ ‖d‖)",
            "Ángulo entre vectores. Robusto a la magnitud; el estándar en embeddings de texto.",
            ACCENT,
        ),
        (
            "Euclídea L2 (curso / lab)",
            "√ Σ (qᵢ − dᵢ)²",
            "Distancia geométrica. Por defecto en la UD3. En AgroPS solo en ablación.",
            ACCENT2,
        ),
        (
            "Manhattan L1 (lab)",
            "Σ |qᵢ − dᵢ|",
            "Suma de desviaciones. Sensible a escala; no usa el índice HNSW.",
            MUTED,
        ),
    ]
    for i, (title, formula, body, color) in enumerate(metrics):
        x = Inches(0.4 + i * 4.25)
        _round(s, x, Inches(1.2), Inches(4.05), Inches(2.55), CARD)
        _rect(s, x, Inches(1.2), Inches(4.05), Inches(0.1), color)
        _textbox(
            s,
            x + Inches(0.18),
            Inches(1.4),
            Inches(3.7),
            Inches(0.4),
            [{"text": title, "size": 15, "bold": True, "color": INK}],
        )
        _textbox(
            s,
            x + Inches(0.18),
            Inches(1.85),
            Inches(3.7),
            Inches(0.45),
            [{"text": formula, "size": 16, "bold": True, "color": color, "font": "Calibri"}],
        )
        _textbox(
            s,
            x + Inches(0.18),
            Inches(2.4),
            Inches(3.7),
            Inches(1.15),
            [{"text": body, "size": 13, "color": MUTED}],
        )

    reasons = [
        ("Por qué coseno", "nomic-embed-text (768-d) se compara por semejanza angular. Padding a 4096 con ceros no cambia el ángulo."),
        ("Elección de k", "k pequeño: ruido (un chunk irrelevante manda). k grande: se diluye la evidencia. AgroPS: retrieval_k = 10 → RRF → final_k."),
        ("Escala / modelos", "k-NN es sensible a escala y a mezclar espacios. Filtramos por model: nunca se compara Nomic con Qwen3."),
        ("Índice ANN", "HNSW (m=16, ef_construction=64) sobre subvector 2000-d. Vecinos aproximados: la predicción deja de ser O(n) brute-force."),
    ]
    for i, (title, body) in enumerate(reasons):
        col, row = i % 2, i // 2
        x = Inches(0.4 + col * 6.4)
        y = Inches(3.95 + row * 1.4)
        _round(s, x, y, Inches(6.15), Inches(1.28), SOFT if i == 0 else CARD)
        _textbox(
            s,
            x + Inches(0.2),
            y + Inches(0.12),
            Inches(5.75),
            Inches(0.32),
            [{"text": title, "size": 14, "bold": True, "color": ACCENT}],
        )
        _textbox(
            s,
            x + Inches(0.2),
            y + Inches(0.48),
            Inches(5.75),
            Inches(0.7),
            [{"text": body, "size": 12, "color": INK}],
        )
    _footer(s, "distancia")

    n = len(prs.slides)
    _move_slide(prs, n - 2, insert_at)
    _move_slide(prs, len(prs.slides) - 1, insert_at + 1)


def _save(prs: Presentation) -> Path:
    try:
        prs.save(str(SRC))
        return SRC
    except OSError as exc:
        prs.save(str(ALT))
        print(f"Defensa abierta o bloqueada ({exc}). Guardado en: {ALT}")
        return ALT


def main() -> int:
    if not SRC.exists():
        print(f"No existe {SRC}", file=sys.stderr)
        return 1
    prs = Presentation(str(SRC))
    if _has_knn(prs):
        print("Las diapositivas de k-NN ya están en el PPTX. No duplico.")
        return 0
    _enrich_embeddings_card(prs)
    insert_at = _find_insert_after(prs)
    add_knn_slides(prs, insert_at)
    out = _save(prs)
    print(f"Guardado: {out} ({len(prs.slides)} diapositivas). k-NN insertado en posición {insert_at + 1}.")
    print("Cierra PowerPoint y vuelve a abrir el archivo si tenías Defensa_TFM_AgroPS_15min.pptx abierto.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
