#!/usr/bin/env python3
"""Presentación de defensa TFM AgroPS (15 min), alineada con GUION_15MIN.md."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
FIGS = ROOT / "docs" / "memoria-latex" / "figures"
OUT = Path(__file__).resolve().parent / "Defensa_TFM_AgroPS_15min.pptx"

TITLE = "Impacto de la Arquitectura RAG en las Alucinaciones de los LLMs"

BG = RGBColor(0xF7, 0xF5, 0xF0)
INK = RGBColor(0x1A, 0x2E, 0x1A)
MUTED = RGBColor(0x4A, 0x5C, 0x4A)
ACCENT = RGBColor(0x2F, 0x6B, 0x3A)
ACCENT2 = RGBColor(0xC4, 0x7A, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xE8, 0xEF, 0xE6)
DANGER = RGBColor(0x8B, 0x2E, 0x2E)
OK = RGBColor(0x1F, 0x6B, 0x3A)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TOTAL = 11


def _set_run(run, text, size=18, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _fill(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    _fill(shape, color)
    return shape


def _round(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    _fill(shape, color)
    return shape


def _textbox(slide, left, top, width, height, paragraphs, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[valign],
        )
    except Exception:
        pass
    first = True
    for pspec in paragraphs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = pspec.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(pspec.get("space_after", 6))
        p.space_before = Pt(pspec.get("space_before", 0))
        run = p.add_run()
        _set_run(
            run,
            pspec["text"],
            size=pspec.get("size", 18),
            bold=pspec.get("bold", False),
            color=pspec.get("color", INK),
            font=pspec.get("font", "Calibri"),
        )
    return box


def _bullet(slide, left, top, width, height, items, size=16, color=INK):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = item.get("level", 0)
        p.space_after = Pt(item.get("space_after", 8))
        run = p.add_run()
        prefix = "• " if p.level == 0 else "– "
        _set_run(
            run,
            prefix + item["text"],
            size=size,
            bold=item.get("bold", False),
            color=color,
        )
    return box


def _badge(slide, left, top, width, height, text, bg=ACCENT, fg=WHITE, size=12):
    shape = _round(slide, left, top, width, height, bg)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    _set_run(run, text, size=size, bold=True, color=fg)
    try:
        tf._txBody.bodyPr.set("anchor", "ctr")
    except Exception:
        pass
    return shape


def _add_picture_safe(slide, path: Path, left, top, width=None, height=None):
    if not path.exists():
        return None
    kwargs = {}
    if width is not None:
        kwargs["width"] = width
    if height is not None:
        kwargs["height"] = height
    return slide.shapes.add_picture(str(path), left, top, **kwargs)


def _footer(slide, n, clock=""):
    _textbox(
        slide,
        Inches(0.4),
        Inches(7.08),
        Inches(10.6),
        Inches(0.32),
        [{"text": TITLE, "size": 10, "color": MUTED}],
    )
    label = f"{n}/{TOTAL}" if not clock else f"{clock}  ·  {n}/{TOTAL}"
    _textbox(
        slide,
        Inches(10.6),
        Inches(7.08),
        Inches(2.4),
        Inches(0.32),
        [{"text": label, "size": 10, "color": MUTED, "align": PP_ALIGN.RIGHT}],
    )


def _title_bar(slide, title, subtitle=None):
    _rect(slide, 0, 0, SLIDE_W, Inches(0.08), ACCENT)
    _textbox(
        slide,
        Inches(0.5),
        Inches(0.22),
        Inches(12.3),
        Inches(0.5),
        [{"text": title, "size": 26, "bold": True, "color": INK}],
    )
    if subtitle:
        _textbox(
            slide,
            Inches(0.5),
            Inches(0.72),
            Inches(12.3),
            Inches(0.32),
            [{"text": subtitle, "size": 14, "color": MUTED}],
        )


def _notes(slide, text: str):
    slide.notes_slide.notes_text_frame.text = text


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, SLIDE_W, SLIDE_H, BG)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = TITLE
    prs.core_properties.author = "Samuel Lorenzo Sánchez"
    prs.core_properties.subject = "Defensa TFM AgroPS · 15 minutos"

    # 1. Portada 0:00–0:40
    s = new_slide(prs)
    _rect(s, 0, 0, Inches(0.35), SLIDE_H, ACCENT)
    _textbox(
        s,
        Inches(0.8),
        Inches(1.45),
        Inches(11.5),
        Inches(0.35),
        [
            {
                "text": "TRABAJO DE FIN DE MÁSTER  ·  INTELIGENCIA ARTIFICIAL  ·  UAX",
                "size": 13,
                "bold": True,
                "color": ACCENT,
            }
        ],
    )
    _textbox(
        s,
        Inches(0.8),
        Inches(1.95),
        Inches(11.6),
        Inches(1.7),
        [{"text": TITLE.replace(" en las ", "\nen las "), "size": 32, "bold": True, "color": INK}],
    )
    _textbox(
        s,
        Inches(0.8),
        Inches(3.85),
        Inches(11.2),
        Inches(0.9),
        [
            {
                "text": "El problema no es la IA en general: es si RAG mitiga alucinaciones\nen documentación agrícola y normativa (PAC / POSEI).",
                "size": 18,
                "color": MUTED,
            }
        ],
    )
    _textbox(
        s,
        Inches(0.8),
        Inches(5.15),
        Inches(8.2),
        Inches(1.2),
        [
            {"text": "Samuel Lorenzo Sánchez", "size": 16, "bold": True, "color": INK},
            {"text": "Tutor: Carlos Granados Aguilar", "size": 14, "color": MUTED},
            {"text": "AgroPS  ·  prototipo contenedorizado  ·  15 minutos", "size": 13, "color": MUTED},
        ],
    )
    _badge(s, Inches(10.15), Inches(5.35), Inches(2.45), Inches(0.48), "15 minutos", ACCENT2)
    _footer(s, 1, "0:00–0:40")
    _notes(
        s,
        "0:00–0:40. Presentarte. Leer el título literal (el mismo que el pie). "
        "Frase: «El problema no es la IA en general: es si RAG mitiga alucinaciones en normativa agraria.»",
    )

    # 2. Pregunta 0:40–2:00
    s = new_slide(prs)
    _title_bar(s, "Pregunta de investigación", "Hilo conductor · 0:40–2:00")
    _round(s, Inches(0.5), Inches(1.2), Inches(12.3), Inches(1.85), SOFT)
    _textbox(
        s,
        Inches(0.75),
        Inches(1.4),
        Inches(11.8),
        Inches(1.5),
        [
            {
                "text": "¿En qué medida un RAG híbrido textual reduce las alucinaciones factuales\nfrente al mismo LLM sin recuperación, en un dominio PAC / POSEI?",
                "size": 22,
                "bold": True,
                "color": INK,
                "align": PP_ALIGN.CENTER,
            }
        ],
    )
    agenda = [
        ("01", "Problema", "2:00"),
        ("02", "Arquitectura", "3:30"),
        ("03", "Corpus", "5:30"),
        ("04", "Cifras", "7:00"),
        ("05", "Demo", "10:30"),
        ("06", "Límites", "12:30"),
    ]
    for i, (num, title, t) in enumerate(agenda):
        x = Inches(0.5 + i * 2.12)
        _round(s, x, Inches(3.35), Inches(1.98), Inches(2.55), CARD)
        _textbox(
            s,
            x + Inches(0.08),
            Inches(3.5),
            Inches(1.82),
            Inches(0.4),
            [{"text": num, "size": 14, "bold": True, "color": ACCENT, "align": PP_ALIGN.CENTER}],
        )
        _textbox(
            s,
            x + Inches(0.08),
            Inches(4.0),
            Inches(1.82),
            Inches(0.85),
            [{"text": title, "size": 16, "bold": True, "color": INK, "align": PP_ALIGN.CENTER}],
        )
        _textbox(
            s,
            x + Inches(0.08),
            Inches(4.95),
            Inches(1.82),
            Inches(0.4),
            [{"text": t, "size": 12, "color": MUTED, "align": PP_ALIGN.CENTER}],
        )
    _footer(s, 2, "0:40–2:00")
    _notes(
        s,
        "Leer la pregunta en voz alta. Agenda: problema → arquitectura → corpus → cifras → demo. "
        "Dejar claro que se compara el mismo modelo con y sin recuperación.",
    )

    # 3. Problema 2:00–3:30
    s = new_slide(prs)
    _title_bar(
        s,
        "El problema: fluido, pero no auditable",
        "Un asesor no puede fiarse de un chat sin fuentes  ·  2:00–3:30",
    )
    _round(s, Inches(0.5), Inches(1.2), Inches(6.15), Inches(2.55), CARD)
    _rect(s, Inches(0.5), Inches(1.2), Inches(0.12), Inches(2.55), DANGER)
    _textbox(s, Inches(0.85), Inches(1.35), Inches(5.55), Inches(0.4), [{"text": "Alucinación extrínseca", "size": 18, "bold": True, "color": DANGER}])
    _textbox(
        s,
        Inches(0.85),
        Inches(1.85),
        Inches(5.55),
        Inches(1.6),
        [
            {
                "text": "El LLM inventa importes, plazos, códigos de ayuda o citas que no están en el expediente. Suena a boletín; no lo es.",
                "size": 16,
                "color": INK,
            }
        ],
    )
    _round(s, Inches(6.85), Inches(1.2), Inches(6.0), Inches(2.55), CARD)
    _rect(s, Inches(6.85), Inches(1.2), Inches(0.12), Inches(2.55), ACCENT2)
    _textbox(s, Inches(7.2), Inches(1.35), Inches(5.4), Inches(0.4), [{"text": "Alucinación intrínseca", "size": 18, "bold": True, "color": ACCENT2}])
    _textbox(
        s,
        Inches(7.2),
        Inches(1.85),
        Inches(5.4),
        Inches(1.6),
        [
            {
                "text": "Contradice el contexto que ya tiene delante. En normativa agraria el fallo habitual es el extrínseco: rellenar huecos con patrones.",
                "size": 16,
                "color": INK,
            }
        ],
    )
    three = [
        ("Sin evidencia", "Pregunta → LLM → texto fluido. No hay fragmento que auditar."),
        ("Sin citas", "El asesor no sabe documento, sección ni chunk."),
        ("Peor que callar", "Una cifra plausible y falsa es más peligrosa que abstenerse."),
    ]
    for i, (t, d) in enumerate(three):
        x = Inches(0.5 + i * 4.2)
        _round(s, x, Inches(4.0), Inches(4.0), Inches(2.35), SOFT)
        _textbox(s, x + Inches(0.25), Inches(4.15), Inches(3.5), Inches(0.45), [{"text": t, "size": 16, "bold": True, "color": INK}])
        _textbox(s, x + Inches(0.25), Inches(4.65), Inches(3.5), Inches(1.45), [{"text": d, "size": 14, "color": MUTED}])
    _footer(s, 3, "2:00–3:30")
    _notes(
        s,
        "LLM fluido inventa importes/plazos/códigos. Distinguir extrínseca vs intrínseca. "
        "Un asesor no puede auditar un chat sin fuentes.",
    )

    # 4. Arquitectura 3:30–5:30
    s = new_slide(prs)
    _title_bar(
        s,
        "Arquitectura: qué se evalúa y qué no",
        "RAG textual modular  ·  citas: documento + sección + chunk  ·  3:30–5:30",
    )
    _add_picture_safe(s, FIGS / "diagrama-respuesta-rag.png", Inches(0.35), Inches(1.12), width=Inches(8.15))
    scopes = [
        (OK, "Evaluado", "RAG textual modular\nDenso + BM25 + RRF\nllama3.2 + nomic-embed"),
        (ACCENT2, "Producto, no corrida", "LangGraph (agentic)\nOCR de ingesta\nLabs de evaluación"),
        (DANGER, "Fuera de alcance", "GraphRAG\nVisión end-to-end\nRAGAS oficial\nCertificación"),
    ]
    for i, (color, title, body) in enumerate(scopes):
        y = Inches(1.12 + i * 1.82)
        _round(s, Inches(8.65), y, Inches(4.25), Inches(1.7), CARD)
        _rect(s, Inches(8.65), y, Inches(0.12), Inches(1.7), color)
        _textbox(s, Inches(8.95), y + Inches(0.12), Inches(3.75), Inches(0.38), [{"text": title, "size": 15, "bold": True, "color": color}])
        _textbox(s, Inches(8.95), y + Inches(0.5), Inches(3.75), Inches(1.1), [{"text": body, "size": 13, "color": MUTED}])
    _footer(s, 4, "3:30–5:30")
    _notes(
        s,
        "Evaluado: RAG textual modular (denso + BM25 + RRF). Producto, no corrida: LangGraph. "
        "Fuera: GraphRAG, visión end-to-end, RAGAS oficial, certificación. "
        "Citas: documento + sección + chunk. Si preguntan por qué híbrido: denso = paráfrasis, BM25 = códigos, RRF fusiona.",
    )

    # 5. Corpus 5:30–7:00
    s = new_slide(prs)
    _title_bar(
        s,
        "Corpus y banco de oro",
        "29 markdown asesor-canarias  ·  N=35 (7×5)  ·  5:30–7:00",
    )
    _round(s, Inches(0.45), Inches(1.18), Inches(6.35), Inches(5.55), CARD)
    _textbox(s, Inches(0.7), Inches(1.35), Inches(5.9), Inches(0.4), [{"text": "Qué se indexó", "size": 16, "bold": True, "color": ACCENT}])
    _bullet(
        s,
        Inches(0.7),
        Inches(1.85),
        Inches(5.9),
        Inches(4.55),
        [
            {"text": "14 ficheros de marco (00–13): POSEI, GIP, filoxera, Panamá, REAC"},
            {"text": "15 dossiers de cultivo (platanera, vid, papa, aguacate…)"},
            {"text": "Destilados de consolidados POSEI 21-01-2025 y 2026 (rige 1-1-2026)"},
            {"text": "1289 fragmentos en la corrida congelada"},
            {"text": "Criterio: vigencia 2023–2026 y anclaje a una sección"},
            {"text": "No se evalúa el censo de beneficiarios ni el JSON de 50 MB"},
        ],
        size=14,
    )
    _round(s, Inches(7.0), Inches(1.18), Inches(5.85), Inches(5.55), SOFT)
    _textbox(s, Inches(7.25), Inches(1.35), Inches(5.4), Inches(0.4), [{"text": "Banco gold_tests.jsonl", "size": 16, "bold": True, "color": ACCENT}])
    examples = [
        ("POSEI = ¿qué sigla?", "02_posei_marco.md"),
        ("243,96 M€ en 2026", "02_posei_marco.md"),
        ("REAC vs CUE vs REA", "10_cuaderno_reac.md"),
        ("Fusarium y el 70 % IGP", "07_mal_de_panama.md"),
        ("AGROIL n.º 22319", "cultivos/platanera.md"),
    ]
    for i, (q, src) in enumerate(examples):
        y = Inches(1.85 + i * 0.72)
        _round(s, Inches(7.25), y, Inches(5.35), Inches(0.64), CARD)
        _textbox(s, Inches(7.4), y + Inches(0.05), Inches(5.05), Inches(0.28), [{"text": q, "size": 13, "bold": True, "color": INK}])
        _textbox(s, Inches(7.4), y + Inches(0.3), Inches(5.05), Inches(0.28), [{"text": src, "size": 11, "color": MUTED}])
    _footer(s, 5, "5:30–7:00")
    _notes(
        s,
        "29 markdown asesor-canarias, destilados de consolidados POSEI 2025–2026. "
        "N=35 (siete categorías × cinco) con respuesta esperada y source_file. No mezclar con un run sintético de 150.",
    )

    # 6. Resultados 7:00–10:30 (números)
    s = new_slide(prs)
    _title_bar(
        s,
        "Resultados: mismo LLM, con y sin RAG",
        "run_20260819T165515Z  ·  N=35  ·  juez local 1–5  ·  7:00–10:30",
    )
    metrics = [
        ("Alucinación (juez)", "31,4 %  →  8,6 %", "Mitiga, no es cero", OK),
        ("Exactitud (1–5)", "3,31  →  4,46", "+1,14 puntos", OK),
        ("MRR híbrido", "0,81", "Primer útil pronto", ACCENT),
        ("Cobertura keywords", "0,96", "Proxy de context recall", ACCENT),
    ]
    for i, (name, val, tag, color) in enumerate(metrics):
        x = Inches(0.45 + (i % 4) * 3.2)
        _round(s, x, Inches(1.2), Inches(3.05), Inches(2.55), CARD)
        _textbox(s, x + Inches(0.18), Inches(1.35), Inches(2.7), Inches(0.55), [{"text": name, "size": 13, "color": MUTED}])
        _textbox(s, x + Inches(0.18), Inches(1.9), Inches(2.7), Inches(1.0), [{"text": val, "size": 22, "bold": True, "color": color}])
        _textbox(s, x + Inches(0.18), Inches(3.05), Inches(2.7), Inches(0.45), [{"text": tag, "size": 12, "color": MUTED}])
    _add_picture_safe(s, FIGS / "metricas-mrr-categorias.png", Inches(0.45), Inches(3.9), width=Inches(7.6))
    _round(s, Inches(8.2), Inches(3.9), Inches(4.7), Inches(2.85), SOFT)
    _textbox(s, Inches(8.4), Inches(4.05), Inches(4.35), Inches(0.4), [{"text": "Cómo leerlo", "size": 15, "bold": True, "color": ACCENT}])
    _bullet(
        s,
        Inches(8.4),
        Inches(4.5),
        Inches(4.35),
        Inches(2.05),
        [
            {"text": "Pareado: mismo modelo, misma pregunta"},
            {"text": "Fortaleza: spanning / compliance"},
            {"text": "Hueco: traceability (MRR 0,67)"},
            {"text": "El juez no es un inspector PAC"},
        ],
        size=13,
    )
    _footer(s, 6, "7:00–10:30")
    _notes(
        s,
        "N=35, run_20260819. Hall. 31,4% → 8,6%. Acc. 3,31 → 4,46. MRR 0,81. "
        "Insistir: no es cero. No citar el gráfico de 21 ítems. No mezclar N=150.",
    )

    # 7. Caso AGROIL (sigue resultados)
    s = new_slide(prs)
    _title_bar(
        s,
        "El 8,6 % residual: trazabilidad",
        "Ítem 34  ·  AGROIL en platanera  ·  el sistema cita cerca e inventa el código",
    )
    _round(s, Inches(0.45), Inches(1.2), Inches(6.2), Inches(5.5), CARD)
    _textbox(s, Inches(0.7), Inches(1.4), Inches(5.75), Inches(0.4), [{"text": "Qué se preguntó", "size": 16, "bold": True, "color": ACCENT}])
    _bullet(
        s,
        Inches(0.7),
        Inches(1.9),
        Inches(5.75),
        Inches(4.5),
        [
            {"text": "Pregunta: n.º de registro de AGROIL para araña roja en platanera", "bold": True},
            {"text": "Referencia de oro: 22319  (cultivos/platanera.md)"},
            {"text": "El dossier entra en el top-k; el chunk con 22319 no"},
            {"text": "Respuesta RAG: inventa ES-01744"},
            {"text": "El juez puso exactitud 5: ve la referencia, no el dígito"},
            {"text": "Lección: citar documento vecino ≠ citar el fragmento correcto"},
        ],
        size=15,
    )
    _round(s, Inches(6.9), Inches(1.2), Inches(5.95), Inches(5.5), SOFT)
    _textbox(s, Inches(7.15), Inches(1.4), Inches(5.5), Inches(0.4), [{"text": "Por eso no prometemos cero", "size": 16, "bold": True, "color": DANGER}])
    _textbox(
        s,
        Inches(7.15),
        Inches(2.0),
        Inches(5.5),
        Inches(4.3),
        [
            {
                "text": "RAG desplaza el riesgo al retriever: si el identificador no está en el contexto, el generador completa con un código plausible.",
                "size": 16,
                "color": INK,
                "space_after": 14,
            },
            {
                "text": "La demo viva repetirá esta pregunta a propósito: es el límite del prototipo, no un accidente que escondemos.",
                "size": 16,
                "color": MUTED,
            },
        ],
    )
    _footer(s, 7, "7:00–10:30")
    _notes(
        s,
        "Hueco: trazabilidad. AGROIL inventa ES-01744 en lugar de 22319. "
        "Juez ≠ humano (puntuó 5). Preparar la misma pregunta en la demo.",
    )

    # 8. Demo 10:30–12:30
    s = new_slide(prs)
    _title_bar(
        s,
        "Demo (90–120 s)  ·  o vídeo si Ollama tarda",
        "Dos preguntas: una fácil de anclar y una que enseña el límite  ·  10:30–12:30",
    )
    _add_picture_safe(s, FIGS / "diagrama-ingesta-rag.png", Inches(0.4), Inches(1.15), width=Inches(6.5))
    _round(s, Inches(7.1), Inches(1.15), Inches(5.75), Inches(5.55), CARD)
    _textbox(s, Inches(7.35), Inches(1.35), Inches(5.3), Inches(0.4), [{"text": "En vivo, en este orden", "size": 16, "bold": True, "color": ACCENT}])
    _round(s, Inches(7.35), Inches(1.9), Inches(5.3), Inches(1.85), SOFT)
    _textbox(s, Inches(7.5), Inches(2.05), Inches(5.0), Inches(0.35), [{"text": "1.  ¿Qué significa POSEI?", "size": 16, "bold": True, "color": INK}])
    _textbox(
        s,
        Inches(7.5),
        Inches(2.45),
        Inches(5.0),
        Inches(1.1),
        [
            {
                "text": "Debe citar 02_posei_marco.md. Señalar documento, sección y fragmento en el panel de fuentes.",
                "size": 14,
                "color": MUTED,
            }
        ],
    )
    _round(s, Inches(7.35), Inches(3.95), Inches(5.3), Inches(2.45), SOFT)
    _textbox(s, Inches(7.5), Inches(4.1), Inches(5.0), Inches(0.55), [{"text": "2.  ¿N.º de registro de AGROIL\nen platanera?", "size": 16, "bold": True, "color": DANGER}])
    _textbox(
        s,
        Inches(7.5),
        Inches(4.75),
        Inches(5.0),
        Inches(1.4),
        [
            {
                "text": "Oro: 22319. Si inventa un código (p. ej. ES-01744), mostrarlo: el retriever falló el chunk y el generador no se abstuvo.",
                "size": 14,
                "color": MUTED,
            }
        ],
    )
    _footer(s, 8, "10:30–12:30")
    _notes(
        s,
        "Chat con fuentes. Si el retrieve tarda, vídeo de 60–90 s. "
        "Pregunta 1 POSEI (éxito). Pregunta 2 AGROIL (límite de trazabilidad). No improvisar una tercera.",
    )

    # 9. Límites 12:30–14:20
    s = new_slide(prs)
    _title_bar(s, "Límites honestos", "Mitigar ≠ certificar  ·  12:30–14:20")
    limits = [
        ("N = 35", "Evidencia principal acotada. Siete categorías × cinco. No es un universo MAPA/FEGA."),
        ("Juez = mismo LLM", "Genera y puntúa. Hay leakage: ve la referencia. No sustituye a un inspector."),
        ("Página PDF", "La cita arrastra source_file, sección y chunk. La página exacta depende del parser."),
        ("Sin ablación", "La corrida mide el híbrido fijo. No dens vs BM25 vs reranker vs agente."),
        ("RAGAS no corrido", "Está en el estado del arte. No hay informe RAGAS oficial en run_20260819."),
        ("No es producción", "Prototipo contenedorizado. No hay certificación de cumplimiento normativo."),
    ]
    for i, (t, d) in enumerate(limits):
        col, row = i % 3, i // 3
        x = Inches(0.45 + col * 4.25)
        y = Inches(1.25 + row * 2.7)
        _round(s, x, y, Inches(4.05), Inches(2.5), CARD)
        _rect(s, x, y, Inches(4.05), Inches(0.1), ACCENT2 if row == 0 else MUTED)
        _textbox(s, x + Inches(0.2), y + Inches(0.3), Inches(3.65), Inches(0.5), [{"text": t, "size": 16, "bold": True, "color": INK}])
        _textbox(s, x + Inches(0.2), y + Inches(0.9), Inches(3.65), Inches(1.35), [{"text": d, "size": 14, "color": MUTED}])
    _footer(s, 9, "12:30–14:20")
    _notes(
        s,
        "N pequeño; mismo modelo genera y juzga; página PDF no siempre en el chunk. Mitigar ≠ certificar.",
    )

    # 10. Cierre 14:20–15:00
    s = new_slide(prs)
    _rect(s, 0, 0, Inches(0.35), SLIDE_H, ACCENT)
    _textbox(
        s,
        Inches(0.85),
        Inches(1.35),
        Inches(11.5),
        Inches(0.4),
        [{"text": "Frase de defensa", "size": 14, "bold": True, "color": ACCENT}],
    )
    _textbox(
        s,
        Inches(0.85),
        Inches(1.85),
        Inches(11.6),
        Inches(2.4),
        [
            {
                "text": "No prometemos cero alucinaciones. Demostramos un anclaje medible a evidencia recuperada, con abstención, citas (documento, sección, fragmento) y una comparación pareada LLM sin RAG frente a RAG.",
                "size": 24,
                "bold": True,
                "color": INK,
            }
        ],
    )
    _round(s, Inches(0.85), Inches(4.55), Inches(11.6), Inches(1.85), SOFT)
    _textbox(
        s,
        Inches(1.15),
        Inches(4.8),
        Inches(11.0),
        Inches(1.4),
        [
            {"text": "Samuel Lorenzo Sánchez  ·  tutor Carlos Granados Aguilar", "size": 16, "bold": True, "color": INK},
            {"text": TITLE, "size": 14, "color": MUTED},
            {"text": "Preguntas", "size": 16, "bold": True, "color": ACCENT},
        ],
    )
    _footer(s, 10, "14:20–15:00")
    _notes(s, "Leer la frase de defensa despacio. Agradecer. Abrir a preguntas.")

    # 11. Si preguntan (reserva)
    s = new_slide(prs)
    _title_bar(s, "Si preguntan", "Diapositiva de reserva  ·  no contar en los 15 min")
    qa = [
        ("¿Por qué híbrido?", "El denso cubre paráfrasis («prima de arranque»). BM25 clava códigos y normas. RRF fusiona sin entrenar un ranker."),
        ("¿GraphRAG / multimodal?", "Fuera del objeto experimental. El OCR solo alimenta texto. UMAP no es un grafo de consulta."),
        ("¿RAGAS?", "Definido en el estado del arte (faithfulness, relevancy, context P/R). No se ejecutó en la corrida congelada."),
        ("¿El run de 150?", "No mezclar. La evidencia canónica es el oro N=35 (run_20260819). El 150 es sensibilidad, no el resultado de defensa."),
    ]
    for i, (q, a) in enumerate(qa):
        y = Inches(1.2 + i * 1.38)
        _round(s, Inches(0.5), y, Inches(12.3), Inches(1.25), CARD)
        _textbox(s, Inches(0.8), y + Inches(0.12), Inches(11.7), Inches(0.35), [{"text": q, "size": 16, "bold": True, "color": ACCENT}])
        _textbox(s, Inches(0.8), y + Inches(0.52), Inches(11.7), Inches(0.55), [{"text": a, "size": 14, "color": INK}])
    _footer(s, 11, "reserva")
    _notes(s, "Solo si el tribunal pregunta. No proyectar en la exposición principal.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"OK -> {OUT}  ({TOTAL} diapositivas)")


if __name__ == "__main__":
    build()
