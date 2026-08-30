#!/usr/bin/env python3
"""Inserta la diapositiva del diagrama embeddings → pgvector → k-NN."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
CANDIDATES = [
    ROOT / "Defensa_TFM_AgroPS_15min_knn.pptx",
    ROOT / "Defensa_TFM_AgroPS_15min.pptx",
]
OUT_FALLBACK = ROOT / "Defensa_TFM_AgroPS_15min_embeddings.pptx"
FIG = ROOT / "diagrama-embeddings-knn.png"

BG = RGBColor(0xF7, 0xF5, 0xF0)
INK = RGBColor(0x1A, 0x2E, 0x1A)
MUTED = RGBColor(0x4A, 0x5C, 0x4A)
ACCENT = RGBColor(0x2F, 0x6B, 0x3A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARKER = "Cómo entra cada texto al embedding"


def _fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    _fill(shape, color)
    return shape


def _textbox(slide, left, top, width, height, specs):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for spec in specs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = spec.get("align", PP_ALIGN.LEFT)
        run = p.add_run()
        run.text = spec["text"]
        run.font.size = Pt(spec.get("size", 14))
        run.font.bold = spec.get("bold", False)
        run.font.color.rgb = spec.get("color", INK)
        run.font.name = "Calibri"
    return box


def _move_slide(prs: Presentation, old_index: int, new_index: int) -> None:
    sld_id_lst = prs.slides._sldIdLst  # noqa: SLF001
    slides = list(sld_id_lst)
    el = slides[old_index]
    sld_id_lst.remove(el)
    sld_id_lst.insert(new_index, el)


def _find_marker_index(prs: Presentation) -> int | None:
    for i, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            if sh.has_text_frame and MARKER in sh.text_frame.text:
                return i
    return None


def _delete_slide(prs: Presentation, index: int) -> None:
    sld_id_lst = prs.slides._sldIdLst  # noqa: SLF001
    el = list(sld_id_lst)[index]
    r_id = el.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    if r_id:
        try:
            prs.part.drop_rel(r_id)
        except Exception:
            pass
    sld_id_lst.remove(el)


def _find_insert_at(prs: Presentation) -> int:
    for i, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            if sh.has_text_frame and "k-NN: de la clasificación (UD3)" in sh.text_frame.text:
                return i + 1
    for i, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            if sh.has_text_frame and "Técnicas de IA aplicadas" in sh.text_frame.text:
                return i + 1
    return min(8, len(prs.slides))


def add_slide(prs: Presentation, insert_at: int) -> None:
    blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    s = prs.slides.add_slide(blank)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, BG)
    _rect(s, 0, 0, SLIDE_W, Inches(0.08), ACCENT)

    _textbox(
        s,
        Inches(0.4),
        Inches(0.16),
        Inches(12.5),
        Inches(0.4),
        [{"text": MARKER, "size": 24, "bold": True, "color": INK}],
    )
    _textbox(
        s,
        Inches(0.4),
        Inches(0.54),
        Inches(12.5),
        Inches(0.32),
        [
            {
                "text": "Documentos y prompt pasan por el mismo modelo. La query busca en pgvector y devuelve los k vecinos (coseno, k=10).",
                "size": 14,
                "color": MUTED,
            }
        ],
    )

    if FIG.exists():
        s.shapes.add_picture(str(FIG), Inches(0.32), Inches(0.95), width=Inches(12.7))
    else:
        _textbox(
            s,
            Inches(1),
            Inches(3),
            Inches(10),
            Inches(1),
            [{"text": "Falta diagrama-embeddings-knn.png", "size": 16, "color": MUTED}],
        )

    n = len(prs.slides)
    _move_slide(prs, n - 1, insert_at)


def main() -> int:
    src = next((p for p in CANDIDATES if p.exists()), None)
    if src is None:
        print("No hay PPTX de defensa.")
        return 1
    prs = Presentation(str(src))
    existing = _find_marker_index(prs)
    insert_at = existing if existing is not None else _find_insert_at(prs)
    if existing is not None:
        _delete_slide(prs, existing)
        insert_at = existing
    add_slide(prs, insert_at)
    try:
        prs.save(str(src))
        out = src
    except OSError as exc:
        prs.save(str(OUT_FALLBACK))
        print(f"Archivo bloqueado ({exc}). Guardado en {OUT_FALLBACK}")
        out = OUT_FALLBACK
    print(f"Guardado: {out} ({len(prs.slides)} diapositivas).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
