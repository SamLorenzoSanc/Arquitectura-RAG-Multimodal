# -*- coding: utf-8 -*-
"""Actualiza Cap. 5 del TFM Word con métricas run_20260820T185427Z (N=150)."""
from __future__ import annotations

import json
import shutil
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT / "docs/evaluation_runs/run_20260820T185427Z"
SUMMARY = json.loads((RUN / "summary.json").read_text(encoding="utf-8"))
FIG_DIR = ROOT / "docs/evaluation_runs/run_20260820T185427Z/figures"
FIG_DIR.mkdir(parents=True, exist_ok=True)

SRC = ROOT / "docs/_TFM_work.docx"
DST_ONEDRIVE = Path(
    r"c:\Users\Usuario\OneDrive - Universidad Alfonso X el Sabio\TFM_Samuel_Lorenzo_Sanchez.docx"
)
DST_DOCS = ROOT / "docs/TFM_Samuel_Lorenzo_Sanchez_ch5_150.docx"

# Estilo visual sobrio (TFM)
COLORS = {
    "rag": "#1B4F72",
    "vanilla": "#7B8A8B",
    "accent": "#196F3D",
    "warn": "#B9770E",
    "grid": "#D5D8DC",
}


def fmt(x: float | None, nd: int = 2) -> str:
    if x is None:
        return "—"
    return f"{x:.{nd}f}".replace(".", ",")


def set_cell_text(cell, text: str, *, bold: bool = False, center: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def fill_table(table, headers: list[str], rows: list[list[str]]) -> None:
    # resize rows
    needed = 1 + len(rows)
    while len(table.rows) < needed:
        table.add_row()
    while len(table.rows) > needed:
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)
    while len(table.columns) < len(headers):
        table.add_column(Cm(2.5))
    # trim extra columns is harder; assume enough columns exist
    for j, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], h, bold=True, center=True)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            set_cell_text(table.rows[i + 1].cells[j], val, center=(j > 0))


def replace_paragraph_image(paragraph, image_path: Path, width_cm: float = 14.5) -> None:
    # clear existing drawings/runs
    p = paragraph._p
    for child in list(p):
        p.remove(child)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))


def find_para_by_prefix(doc: Document, prefix: str, start: int = 470, end: int = 530):
    for i in range(start, end):
        if doc.paragraphs[i].text.strip().startswith(prefix):
            return doc.paragraphs[i]
    raise KeyError(prefix)


def find_body_image_paras(doc: Document):
    """Return (fig1_para, fig2_para) empty paragraphs that hold drawings in ch5."""
    body = list(doc.element.body)
    found = []
    in_ch5 = False
    for child in body:
        tag = child.tag.split("}")[-1]
        if tag != "p":
            continue
        texts = [
            t.text
            for t in child.findall(
                ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
            )
            if t.text
        ]
        joined = "".join(texts)
        if "Presentación de análisis, resultados" in joined:
            in_ch5 = True
        if in_ch5 and "Capítulo 6:" in joined:
            break
        has_d = bool(child.xpath('.//*[local-name()="drawing"]'))
        if in_ch5 and has_d:
            # map element to paragraph object
            for para in doc.paragraphs:
                if para._element is child:
                    found.append(para)
                    break
    if len(found) < 2:
        raise RuntimeError(f"Se esperaban 2 figuras en cap.5, hay {len(found)}")
    return found[0], found[1]


def make_figures(g: dict, by_cat: dict) -> tuple[Path, Path, Path]:
    # Orden tipológico del banco 150
    order = [
        "direct_fact",
        "temporal",
        "spanning",
        "comparative",
        "numerical",
        "relationship",
        "holistic",
    ]
    labels = {
        "direct_fact": "direct_fact",
        "temporal": "temporal",
        "spanning": "spanning",
        "comparative": "comparative",
        "numerical": "numerical",
        "relationship": "relationship",
        "holistic": "holistic",
    }

    # Fig 1: MRR por categoría
    cats = [c for c in order if c in by_cat]
    mrrs = [by_cat[c]["mrr"] for c in cats]
    fig, ax = plt.subplots(figsize=(10, 4.8), dpi=160)
    y = np.arange(len(cats))
    bars = ax.barh(y, mrrs, color=COLORS["rag"], height=0.65)
    ax.axvline(g["mrr"], color=COLORS["warn"], linestyle="--", linewidth=1.5, label=f"MRR global ({fmt(g['mrr'], 3)})")
    ax.set_yticks(y)
    ax.set_yticklabels([labels[c] for c in cats], fontsize=10)
    ax.set_xlabel("MRR")
    ax.set_xlim(0, 1.05)
    ax.set_title("MRR del recuperador híbrido por categoría (N=150)")
    ax.grid(axis="x", color=COLORS["grid"], linestyle=":", linewidth=0.8)
    ax.legend(loc="lower right", frameon=False)
    for bar, v in zip(bars, mrrs):
        ax.text(v + 0.02, bar.get_y() + bar.get_height() / 2, fmt(v, 3), va="center", fontsize=9)
    fig.tight_layout()
    fig1 = FIG_DIR / "fig_mrr_categorias_150.png"
    fig.savefig(fig1, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    # Fig 2: RAG vs vanilla
    metrics = ["Exactitud", "Exhaustividad", "Relevancia"]
    rag_vals = [g["rag_accuracy"], g["rag_completeness"], g["rag_relevance"]]
    van_vals = [g["vanilla_accuracy"], g["vanilla_completeness"], g["vanilla_relevance"]]
    x = np.arange(len(metrics))
    w = 0.36
    fig, ax = plt.subplots(figsize=(9.5, 5.0), dpi=160)
    b1 = ax.bar(x - w / 2, van_vals, w, label="Sin RAG", color=COLORS["vanilla"])
    b2 = ax.bar(x + w / 2, rag_vals, w, label="Con RAG", color=COLORS["rag"])
    ax.set_xticks(x)
    ax.set_xticklabels(metrics)
    ax.set_ylim(0, 5.4)
    ax.set_ylabel("Nota media LLM-as-a-Judge (1–5)")
    ax.set_title("Comparación sin RAG vs con RAG (N=150, 147 con juez)")
    ax.grid(axis="y", color=COLORS["grid"], linestyle=":", linewidth=0.8)
    ax.legend(frameon=False)
    for bars in (b1, b2):
        for bar in bars:
            h = bar.get_height()
            ax.text(bar.get_x() + bar.get_width() / 2, h + 0.08, fmt(h, 2), ha="center", va="bottom", fontsize=9)
    # hallucination annotation
    ax.text(
        0.02,
        0.98,
        f"Alucinación (juez): sin RAG {fmt(g['vanilla_hallucination_rate'], 3)} → con RAG {fmt(g['rag_hallucination_rate'], 3)}",
        transform=ax.transAxes,
        va="top",
        fontsize=9,
        color=COLORS["accent"],
    )
    fig.tight_layout()
    fig2 = FIG_DIR / "fig_rag_vs_vanilla_150.png"
    fig.savefig(fig2, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    # Fig 3: exactitud RAG vs vanilla por categoría
    acc_r = [by_cat[c]["rag_accuracy"] for c in cats]
    acc_v = [by_cat[c]["vanilla_accuracy"] for c in cats]
    x = np.arange(len(cats))
    w = 0.38
    fig, ax = plt.subplots(figsize=(11, 5.0), dpi=160)
    ax.bar(x - w / 2, acc_v, w, label="Sin RAG", color=COLORS["vanilla"])
    ax.bar(x + w / 2, acc_r, w, label="Con RAG", color=COLORS["rag"])
    ax.set_xticks(x)
    ax.set_xticklabels([labels[c] for c in cats], rotation=20, ha="right")
    ax.set_ylim(0, 5.4)
    ax.set_ylabel("Exactitud media (1–5)")
    ax.set_title("Exactitud del juez por tipología (RAG vs sin RAG)")
    ax.grid(axis="y", color=COLORS["grid"], linestyle=":", linewidth=0.8)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig3 = FIG_DIR / "fig_accuracy_por_categoria_150.png"
    fig.savefig(fig3, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    return fig1, fig2, fig3


def main() -> None:
    g = SUMMARY["global"]
    by = SUMMARY["by_category"]
    fig1, fig2, fig3 = make_figures(g, by)

    doc = Document(str(SRC))

    # --- Textos ---
    intro = (
        "Este capítulo presenta el análisis, los resultados y la discusión del protocolo experimental. "
        "La muestra no es una encuesta sociodemográfica de usuarios, sino un banco de evaluación versionado "
        "(gateway/routes/tests.jsonl) con N=150 preguntas sintéticas sobre el corpus AgroLLM + normativa "
        "agraria canaria (POSEI/GIP). Las cifras proceden de la ejecución congelada "
        "run_20260820T185427Z (150 filas; 147/150 con juez completo; 3 timeouts en los ítems 10, 45 y 57; "
        "duración total ≈ 41,2 min). Misma pregunta sin contexto y con recuperador híbrido "
        "(denso + BM25 + RRF, RAG_USE_RERANKER=false, chunk 1400/120, final_k=8). RAGAS no se ejecutó."
    )
    find_para_by_prefix(doc, "Este capítulo presenta").text = intro

    unidad = (
        "Cada ítem del banco incluye pregunta, respuesta de referencia, palabras clave y categoría tipológica. "
        "El índice se construyó sobre knowledge-base completo (compañía, empleados, productos y asesor-canarias). "
        "La distribución tipológica es desbalanceada a propósito: direct_fact (70), temporal (20), spanning (20), "
        "comparative (10), numerical (10), relationship (10) y holistic (10)."
    )
    find_para_by_prefix(doc, "Cada ítem del banco").text = unidad

    var_rec = (
        "De recuperación: posición del primer fragmento útil (MRR) y cobertura de palabras clave. "
        f"El nDCG del harness no se cita como canónico: produjo valores anómalos >1 en {g['ndcg_gt1']} de "
        f"{SUMMARY['n_ok_no_error']} filas evaluadas."
    )
    find_para_by_prefix(doc, "De recuperación:").text = var_rec

    # IR bullets
    find_para_by_prefix(doc, "MRR global:").text = f"MRR global: {fmt(g['mrr'], 4)}"
    find_para_by_prefix(doc, "Cobertura de keywords:").text = (
        f"Cobertura de keywords: {fmt(g['coverage'], 4)}"
    )

    # best/worst MRR
    ranked = sorted(
        ((k, v["mrr"]) for k, v in by.items() if v.get("mrr") is not None),
        key=lambda kv: kv[1],
        reverse=True,
    )
    best, worst = ranked[0], ranked[-1]
    find_para_by_prefix(doc, "Por categoría, el MRR").text = (
        f"Por categoría, el MRR más alto se observa en {best[0]} ({fmt(best[1], 3)}) "
        f"y el más bajo en {worst[0]} ({fmt(worst[1], 3)})."
    )

    find_para_by_prefix(doc, "Tabla 2. MRR").text = (
        "Tabla 2. MRR del recuperador híbrido por categoría (run_20260820T185427Z)."
    )
    find_para_by_prefix(doc, "Figura 1.").text = (
        "Figura 1. MRR por categoría. Valores numéricos: Tabla 2 (N=150)."
    )

    find_para_by_prefix(doc, "Exactitud:").text = f"Exactitud: {fmt(g['rag_accuracy'], 2)}"
    find_para_by_prefix(doc, "Exhaustividad:").text = (
        f"Exhaustividad: {fmt(g['rag_completeness'], 2)}"
    )
    find_para_by_prefix(doc, "Relevancia:").text = f"Relevancia: {fmt(g['rag_relevance'], 2)}"

    find_para_by_prefix(doc, "Tabla 3. LLM").text = (
        "Tabla 3. LLM autónomo frente a RAG híbrido (N=150; 147 con juez; run_20260820T185427Z)."
    )
    find_para_by_prefix(doc, "Figura 2.").text = (
        "Figura 2. Comparación sin RAG vs con RAG. Valores numéricos: Tabla 3."
    )

    disc = (
        f"El RAG mejora las tres dimensiones del juez (exactitud {fmt(g['vanilla_accuracy'], 2)}→"
        f"{fmt(g['rag_accuracy'], 2)}; exhaustividad {fmt(g['vanilla_completeness'], 2)}→"
        f"{fmt(g['rag_completeness'], 2)}; relevancia {fmt(g['vanilla_relevance'], 2)}→"
        f"{fmt(g['rag_relevance'], 2)}) y reduce la tasa binaria de alucinación de forma "
        f"modesta ({fmt(g['vanilla_hallucination_rate'], 3)} → {fmt(g['rag_hallucination_rate'], 3)}), "
        "sin eliminarla. El mayor salto de exactitud aparece en relationship y holistic; spanning "
        "queda empatado entre con/sin RAG; numerical concentra más alucinación residual (0,30)."
    )
    find_para_by_prefix(doc, "El RAG mejora").text = disc

    find_para_by_prefix(doc, "Tabla 4. MRR").text = (
        "Tabla 4. MRR, exactitud del juez y alucinación por categoría (misma corrida)."
    )
    find_para_by_prefix(doc, "Tabla 5. Origen").text = (
        "Tabla 5. Origen de cada métrica citada en este capítulo."
    )
    # fix typo heading if present
    for i in range(510, 525):
        if doc.paragraphs[i].text.strip().startswith("Origen de las métricas"):
            doc.paragraphs[i].text = "Origen de las métricas reportadas"
            break

    # Insert caption for fig3 before "Origen..." if not present
    origen_idx = None
    for i in range(510, 525):
        if doc.paragraphs[i].text.strip().startswith("Origen de las métricas"):
            origen_idx = i
            break
    # Add figure 3 after table 4 discussion area: look for empty Título Apartado before Origen
    # We'll append fig3 image + caption just before Origen paragraph by editing previous empty paras
    # Find paragraph index of Tabla 4 caption and use following empty slots.
    tabla4_idx = None
    for i in range(505, 525):
        if doc.paragraphs[i].text.strip().startswith("Tabla 4."):
            tabla4_idx = i
            break

    # --- Tables (doc.tables indices from dump) ---
    # Table 4 = config
    fill_table(
        doc.tables[4],
        ["Parámetro", "Valor"],
        [
            ["run_id", "run_20260820T185427Z"],
            ["Banco", "gateway/routes/tests.jsonl (N=150)"],
            ["Generador", "llama3.2:latest"],
            ["Embedder", "nomic-embed-text:latest"],
            ["Chunk / solape", "1400 / 120 caracteres"],
            ["Retrieve", "denso k=10 + BM25 k=10 + RRF (k=60), final_k=8"],
            ["Reranker", "false"],
            ["Temperatura", "0,0 (config)"],
            ["Corpus", "knowledge-base/** (compañía + productos + empleados + asesor-canarias)"],
            ["Filas / juez / errores", "150 / 147 / 3 timeouts (ítems 10, 45, 57)"],
            ["Duración", "≈ 41,2 min (2026-08-20T18:54Z → 19:35Z)"],
            ["RAGAS", "no ejecutado"],
        ],
    )

    # Table 5 = MRR by category
    order = [
        "holistic",
        "direct_fact",
        "relationship",
        "spanning",
        "comparative",
        "temporal",
        "numerical",
    ]
    mrr_rows = []
    for cat in order:
        v = by[cat]
        mrr_rows.append([cat, str(v["n"]), fmt(v["mrr"], 3)])
    mrr_rows.append(["Global", str(SUMMARY["n_ok_no_error"]), fmt(g["mrr"], 3)])
    fill_table(doc.tables[5], ["Categoría", "N", "MRR"], mrr_rows)

    # Table 6 = RAG vs vanilla
    fill_table(
        doc.tables[6],
        ["Métrica", "Sin RAG", "Con RAG", "Dif."],
        [
            [
                "Exactitud (1–5)",
                fmt(g["vanilla_accuracy"], 2),
                fmt(g["rag_accuracy"], 2),
                f"+{fmt(g['rag_accuracy'] - g['vanilla_accuracy'], 2)}",
            ],
            [
                "Exhaustividad (1–5)",
                fmt(g["vanilla_completeness"], 2),
                fmt(g["rag_completeness"], 2),
                f"+{fmt(g['rag_completeness'] - g['vanilla_completeness'], 2)}",
            ],
            [
                "Relevancia (1–5)",
                fmt(g["vanilla_relevance"], 2),
                fmt(g["rag_relevance"], 2),
                f"+{fmt(g['rag_relevance'] - g['vanilla_relevance'], 2)}",
            ],
            [
                "Tasa is_hallucination (juez)",
                fmt(g["vanilla_hallucination_rate"], 3),
                fmt(g["rag_hallucination_rate"], 3),
                fmt(g["rag_hallucination_rate"] - g["vanilla_hallucination_rate"], 3),
            ],
        ],
    )

    # Table 7 = by category detailed
    det_rows = []
    for cat in [
        "direct_fact",
        "temporal",
        "spanning",
        "comparative",
        "numerical",
        "relationship",
        "holistic",
    ]:
        v = by[cat]
        det_rows.append(
            [
                cat,
                str(v["n_judge"]),
                fmt(v["mrr"], 3),
                fmt(v["rag_accuracy"], 2),
                fmt(v["vanilla_accuracy"], 2),
                fmt(v["rag_hallucination_rate"], 2),
            ]
        )
    fill_table(
        doc.tables[7],
        ["Categoría", "N juez", "MRR", "Acc. RAG", "Acc. sin RAG", "Halluc. RAG"],
        det_rows,
    )

    # Table 8 = origen
    fill_table(
        doc.tables[8],
        ["Métrica", "Origen", "Dónde se calcula"],
        [
            [
                "MRR y cobertura de keywords",
                "Función propia",
                "results.csv de run_20260820T185427Z",
            ],
            [
                "nDCG del harness",
                "Función propia, no citada como canónica",
                f"{g['ndcg_gt1']}/{SUMMARY['n_ok_no_error']} filas >1",
            ],
            [
                "Exactitud, exhaustividad, relevancia 1–5, is_hallucination",
                "LLM-as-a-Judge (mismo LLM)",
                "Mismo CSV; no es juicio humano",
            ],
            ["RAGAS", "No aplica", "No se ejecutó"],
            [
                "Timeouts",
                "Ejecución Ollama/juez",
                "Ítems 10, 45 y 57 (timed out)",
            ],
        ],
    )

    # --- Images ---
    fig1_p, fig2_p = find_body_image_paras(doc)
    replace_paragraph_image(fig1_p, fig1)
    replace_paragraph_image(fig2_p, fig2)

    # Figura 3: insertar imagen + pie justo antes de "Origen de las métricas..."
    cap_txt = (
        "Figura 3. Exactitud del juez por tipología (RAG vs sin RAG). "
        "Valores numéricos: Tabla 4."
    )
    if origen_idx is not None:
        # Si ya hay un párrafo vacío inmediatamente antes, reutilizarlo
        prev = doc.paragraphs[origen_idx - 1]
        if prev.text.strip() == "" and not prev._element.xpath(
            './/*[local-name()="drawing"]'
        ):
            replace_paragraph_image(prev, fig3, width_cm=15.0)
            doc.paragraphs[origen_idx].insert_paragraph_before(cap_txt)
        else:
            img_p = doc.paragraphs[origen_idx].insert_paragraph_before("")
            replace_paragraph_image(img_p, fig3, width_cm=15.0)
            # el insert anterior queda detrás; añadir caption antes de Origen de nuevo
            # Recalcular índice Origen
            for i, p in enumerate(doc.paragraphs):
                if p.text.strip().startswith("Origen de las métricas"):
                    p.insert_paragraph_before(cap_txt)
                    break

    doc.save(str(SRC))
    shutil.copy2(SRC, DST_DOCS)
    try:
        shutil.copy2(SRC, DST_ONEDRIVE)
        onedrive_ok = True
    except Exception as exc:  # noqa: BLE001
        onedrive_ok = False
        print("WARN OneDrive copy failed:", exc)

    print("Figures:", fig1.name, fig2.name, fig3.name)
    print("Saved work:", SRC)
    print("Saved docs:", DST_DOCS)
    print("OneDrive:", "OK" if onedrive_ok else "FAILED")


if __name__ == "__main__":
    main()
