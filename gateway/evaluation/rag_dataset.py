"""Persistencia, normalización y conectores del Wizard de datasets RAG."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
from collections.abc import Awaitable, Callable
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable
from uuid import uuid4

from fastapi import HTTPException, UploadFile
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from services.video_extract import (
    MAX_UPLOAD_BYTES,
    is_video_file,
    transcribe_video_segments,
)

RAG_FIELDS = (
    "traceId",
    "prompt",
    "context",
    "response",
    "timestamp",
    "expected_response",
    "expected_context",
    "metadata",
    "pipeline",
    "alternate_response",
    "code_hash",
)
ALIASES = {
    "traceId": ("traceId", "trace_id", "id"),
    "prompt": (
        "prompt",
        "question",
        "query",
        "input",
        "pregunta",
        "topic",
        "headline",
        "user_message",
        "message",
        "log_message",
    ),
    "context": (
        "context",
        "contexts",
        "retrieved_context",
        "contexto",
        "description",
        "fragment",
        "content",
        "chunk",
    ),
    "response": (
        "response",
        "answer",
        "output",
        "respuesta",
        "assistant_message",
        "completion",
    ),
    "timestamp": ("timestamp", "created_at", "date", "fecha"),
    "expected_response": (
        "expected_response",
        "reference_answer",
        "ground_truth",
        "expected_answer",
        "applications",
        "summary",
    ),
    "expected_context": (
        "expected_context",
        "reference_context",
        "ground_truth_context",
    ),
    "metadata": ("metadata", "meta"),
    "pipeline": ("pipeline", "pipeline_config"),
    "alternate_response": ("alternate_response", "alternative_answer"),
    "code_hash": ("code_hash", "hash"),
}
TABULAR_SUFFIXES = {".csv", ".json", ".jsonl", ".ndjson", ".xlsx", ".xls"}
LOG_SUFFIXES = {".log"}
_NUMBER_RE = re.compile(r"^-?\d+(?:[.,]\d+)?$")
_INT_RE = re.compile(r"^-?\d+$")
_TABLES_READY = False
RAG_DATASET_DDL = (
    """CREATE TABLE IF NOT EXISTS public.department_knowledge_bases (
        department_id uuid NOT NULL REFERENCES public.departments(id) ON DELETE CASCADE,
        knowledge_base_id uuid NOT NULL REFERENCES public.knowledge_bases(id) ON DELETE CASCADE,
        organization_id uuid NOT NULL REFERENCES public.organizations(id) ON DELETE CASCADE,
        created_by uuid REFERENCES public.users(id) ON DELETE SET NULL,
        created_at timestamptz NOT NULL DEFAULT CURRENT_TIMESTAMP,
        PRIMARY KEY (department_id, knowledge_base_id))""",
    """CREATE INDEX IF NOT EXISTS idx_department_kb_org
        ON public.department_knowledge_bases (organization_id)""",
    """CREATE TABLE IF NOT EXISTS public.rag_datasets (
        id uuid PRIMARY KEY DEFAULT gen_random_uuid(),
        organization_id uuid NOT NULL REFERENCES public.organizations(id) ON DELETE CASCADE,
        tenant_id uuid NOT NULL REFERENCES public.tenants(id) ON DELETE CASCADE,
        knowledge_base_id uuid REFERENCES public.knowledge_bases(id) ON DELETE SET NULL,
        name varchar(255) NOT NULL, description text,
        source_type varchar(40) NOT NULL DEFAULT 'upload', source_name text,
        source_metadata jsonb NOT NULL DEFAULT '{}'::jsonb,
        mapping jsonb NOT NULL DEFAULT '{}'::jsonb,
        status varchar(30) NOT NULL DEFAULT 'ready',
        row_count integer NOT NULL DEFAULT 0,
        created_by uuid REFERENCES public.users(id) ON DELETE SET NULL,
        created_at timestamptz NOT NULL DEFAULT CURRENT_TIMESTAMP,
        updated_at timestamptz NOT NULL DEFAULT CURRENT_TIMESTAMP)""",
    """CREATE INDEX IF NOT EXISTS idx_rag_datasets_scope
        ON public.rag_datasets (organization_id, tenant_id, created_at DESC)""",
    """CREATE INDEX IF NOT EXISTS idx_rag_datasets_kb
        ON public.rag_datasets (knowledge_base_id)""",
    """CREATE TABLE IF NOT EXISTS public.rag_dataset_rows (
        id uuid PRIMARY KEY DEFAULT gen_random_uuid(),
        dataset_id uuid NOT NULL REFERENCES public.rag_datasets(id) ON DELETE CASCADE,
        organization_id uuid NOT NULL REFERENCES public.organizations(id) ON DELETE CASCADE,
        tenant_id uuid NOT NULL REFERENCES public.tenants(id) ON DELETE CASCADE,
        knowledge_base_id uuid REFERENCES public.knowledge_bases(id) ON DELETE SET NULL,
        source varchar(80), status varchar(30) NOT NULL DEFAULT 'ready',
        "traceId" text, prompt text NOT NULL, context jsonb, response text,
        timestamp timestamptz, expected_response text, expected_context jsonb,
        metadata jsonb NOT NULL DEFAULT '{}'::jsonb, pipeline jsonb,
        alternate_response text, code_hash varchar(64) NOT NULL,
        created_at timestamptz NOT NULL DEFAULT CURRENT_TIMESTAMP,
        UNIQUE (dataset_id, code_hash))""",
    """CREATE INDEX IF NOT EXISTS idx_rag_dataset_rows_scope
        ON public.rag_dataset_rows (organization_id, tenant_id)""",
    """CREATE INDEX IF NOT EXISTS idx_rag_dataset_rows_dataset
        ON public.rag_dataset_rows (dataset_id, created_at)""",
    """CREATE TABLE IF NOT EXISTS public.rag_dataset_departments (
        dataset_id uuid NOT NULL REFERENCES public.rag_datasets(id) ON DELETE CASCADE,
        department_id uuid NOT NULL REFERENCES public.departments(id) ON DELETE CASCADE,
        PRIMARY KEY (dataset_id, department_id))""",
    """CREATE INDEX IF NOT EXISTS idx_rag_dataset_departments_dept
        ON public.rag_dataset_departments (department_id)""",
)


async def init_rag_dataset_tables(db: AsyncSession) -> None:
    """Asegura el DDL del wizard sin leer ficheros fuera de la imagen Docker."""
    global _TABLES_READY
    if _TABLES_READY:
        return
    try:
        exists = await db.scalar(text("SELECT to_regclass('public.rag_datasets')"))
    except Exception:
        return
    if exists:
        _TABLES_READY = True
        return
    for statement in RAG_DATASET_DDL:
        await db.execute(text(statement))
    await db.commit()
    _TABLES_READY = True


def _json_value(value: Any, fallback: Any) -> Any:
    if value is None or value == "":
        return fallback
    if isinstance(value, (dict, list, int, float, bool)):
        return value
    try:
        return json.loads(str(value))
    except (TypeError, json.JSONDecodeError):
        return [str(value)] if isinstance(fallback, list) else {"value": str(value)}


def infer_mapping(columns: Iterable[str]) -> dict[str, str]:
    actual = {str(c).strip().lower(): str(c) for c in columns}
    mapping: dict[str, str] = {}
    for target, aliases in ALIASES.items():
        for alias in aliases:
            if alias.lower() in actual:
                mapping[target] = actual[alias.lower()]
                break
    return mapping


def normalize_row(row: dict[str, Any], mapping: dict[str, str]) -> dict[str, Any]:
    def source(field: str) -> Any:
        key = mapping.get(field)
        return row.get(key) if key else None

    def text_value(field: str) -> str | None:
        value = source(field)
        if isinstance(value, list):
            value = "\n".join(str(item) for item in value if item not in (None, ""))
        return str(value or "").strip() or None

    prompt = str(source("prompt") or "").strip()
    if not prompt:
        fragment = source("context")
        if isinstance(fragment, list):
            fragment = " ".join(str(item) for item in fragment if item)
        prompt = str(fragment or "").strip()[:240]
    if not prompt:
        raise ValueError("La fila no contiene prompt")
    raw_timestamp = source("timestamp")
    parsed_timestamp = None
    if isinstance(raw_timestamp, datetime):
        parsed_timestamp = raw_timestamp
    elif raw_timestamp not in (None, ""):
        try:
            parsed_timestamp = datetime.fromisoformat(
                str(raw_timestamp).strip().replace("Z", "+00:00")
            )
        except ValueError as exc:
            raise ValueError("timestamp no es ISO-8601") from exc
    normalized = {
        "traceId": str(source("traceId") or "").strip() or None,
        "prompt": prompt,
        "context": _json_value(source("context"), []),
        "response": text_value("response"),
        "timestamp": parsed_timestamp,
        "expected_response": text_value("expected_response"),
        "expected_context": _json_value(source("expected_context"), []),
        "metadata": _json_value(source("metadata"), {}),
        "pipeline": _json_value(source("pipeline"), {}),
        "alternate_response": text_value("alternate_response"),
    }
    metadata = normalized["metadata"]
    if not isinstance(metadata, dict):
        metadata = {"value": metadata}
    mapped_keys = {value for value in mapping.values() if value}
    extras = {
        str(key): _cell_value(value)
        for key, value in row.items()
        if str(key) not in mapped_keys
    }
    if extras:
        metadata = {**metadata, "source_columns": extras}
    normalized["metadata"] = metadata
    canonical = json.dumps(normalized, sort_keys=True, ensure_ascii=False, default=str)
    supplied_hash = str(source("code_hash") or "").strip()
    normalized["code_hash"] = (
        supplied_hash or hashlib.sha256(canonical.encode("utf-8")).hexdigest()
    )
    return normalized


def _parse_log_lines(decoded: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for line in decoded.splitlines():
        raw = line.strip()
        if not raw:
            continue
        try:
            payload = json.loads(raw)
            if isinstance(payload, dict):
                rows.append(payload)
                continue
        except json.JSONDecodeError:
            pass
        rows.append({"prompt": raw, "metadata": {"origin": "log_line"}})
    if not rows:
        raise ValueError("El log no contiene líneas importables")
    return rows


DEMO_CATALOG: list[dict[str, Any]] = [
    {
        "id": "posei-faq",
        "name": "FAQ POSEI Canarias",
        "description": "Preguntas frecuentes sobre ayudas POSEI y justificación de pagos.",
        "rows": [
            {
                "prompt": "¿Qué es el programa POSEI en Canarias?",
                "expected_response": (
                    "El POSEI es el régimen de ayudas de la UE para regiones ultraperiféricas. "
                    "En Canarias compensa costes extra de producción y comercialización agrícola."
                ),
                "context": [
                    "POSEI: Programa de Opciones Específicas por la Lejanía y la Insularidad."
                ],
            },
            {
                "prompt": "¿Qué documentos se piden para justificar una ayuda POSEI?",
                "expected_response": (
                    "Normalmente facturas, albaranes, certificados de entrega y el cuaderno de "
                    "explotación. Conserva los justificantes el plazo que indique la convocatoria."
                ),
                "context": ["Justificación: facturas, albaranes y registro de explotación."],
            },
            {
                "prompt": "¿Puedo pedir POSEI si cultivo plátano y tomate en la misma finca?",
                "expected_response": (
                    "Sí, si cada cultivo cumple su ficha y superficies. Declara cada línea por "
                    "separado y no mezcles superficies en el mismo expediente."
                ),
                "context": ["Un expediente puede incluir varias líneas de cultivo."],
            },
        ],
    },
    {
        "id": "riego-gotero",
        "name": "Riego por goteo — casos de campo",
        "description": "Diagnóstico y recambios habituales de goteros, filtros y presión.",
        "rows": [
            {
                "prompt": "Los goteros de una línea no riegan igual. ¿Qué reviso primero?",
                "expected_response": (
                    "Revisa filtro de malla, presión a pie de ramal y taponamiento. Limpia el "
                    "filtro, abre finales de línea y sustituye goteros ciegos."
                ),
                "context": ["Causa habitual: filtro sucio o gotero taponado por cal."],
            },
            {
                "prompt": "¿Qué recambios llevo para un fallo de riego por goteo?",
                "expected_response": (
                    "Goteros, juntas, filtro de malla, llave de paso y cinta de teflón. "
                    "Si hay arena, lleva también un filtro de anillas de recambio."
                ),
                "context": ["Recambios: goteros, filtros, juntas y llaves."],
            },
        ],
    },
    {
        "id": "sanidad-platanera",
        "name": "Sanidad vegetal — platanera",
        "description": "Trips, sigatoka y clorosis: preguntas de diagnóstico rápido.",
        "rows": [
            {
                "prompt": "Veo plateado en el envés de la hoja del plátano. ¿Qué puede ser?",
                "expected_response": (
                    "Es compatible con trips. Confirma con lupa, retira hojas muy afectadas y "
                    "consulta el tratamiento autorizado para platanera en tu zona."
                ),
                "context": ["Trips: plateado y punteado en envés. No improvises materia activa."],
            },
            {
                "prompt": "Manchas foliares alargadas en platanera: ¿sigatoka?",
                "expected_response": (
                    "Puede ser sigatoka. Mejora aireación, deshoja material enfermo y sigue el "
                    "calendario de Sanidad Vegetal. No uses un fungicida sin ficha autorizada."
                ),
                "context": ["Sigatoka: manchas foliares; manejo cultural y productos autorizados."],
            },
        ],
    },
]


def demo_catalog() -> list[dict[str, Any]]:
    return [
        {
            "id": item["id"],
            "name": item["name"],
            "description": item["description"],
            "row_count": len(item["rows"]),
        }
        for item in DEMO_CATALOG
    ]


def demo_rows(catalog_id: str) -> tuple[str, str, list[dict[str, Any]]]:
    match = next((item for item in DEMO_CATALOG if item["id"] == catalog_id), None)
    if match is None:
        raise HTTPException(status_code=404, detail="Dataset demo no encontrado")
    return match["name"], match["description"], list(match["rows"])


def _cell_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="ignore")
    return value


def _unique_names(names: Iterable[Any]) -> list[str]:
    seen: dict[str, int] = {}
    result: list[str] = []
    for index, raw in enumerate(names):
        name = str(raw).strip() if raw not in (None, "") else f"col_{index + 1}"
        count = seen.get(name, 0) + 1
        seen[name] = count
        result.append(name if count == 1 else f"{name}_{count}")
    return result


def infer_column_type(values: Iterable[Any]) -> str:
    samples = [value for value in values if value not in (None, "")]
    if not samples:
        return "string"
    probe = samples[:40]
    if all(
        isinstance(value, bool)
        or str(value).strip().lower() in {"true", "false"}
        for value in probe
    ) and not all(_INT_RE.match(str(value).strip()) for value in probe):
        return "bool"
    if all(isinstance(value, (dict, list)) for value in probe):
        return "json"
    if all(isinstance(value, datetime) or isinstance(value, date) for value in probe):
        return "datetime"
    ints = 0
    floats = 0
    dates = 0
    for value in probe:
        if isinstance(value, bool):
            continue
        if isinstance(value, int):
            ints += 1
            continue
        if isinstance(value, float):
            floats += 1
            continue
        text = str(value).strip()
        if _INT_RE.match(text):
            ints += 1
            continue
        if _NUMBER_RE.match(text.replace(" ", "")):
            floats += 1
            continue
        try:
            datetime.fromisoformat(text.replace("Z", "+00:00"))
            dates += 1
        except ValueError:
            pass
    total = len(probe)
    if dates == total:
        return "datetime"
    if ints == total:
        return "int"
    if ints + floats == total:
        return "float"
    return "string"


def generate_schema(
    columns: Iterable[str], rows: list[dict[str, Any]]
) -> dict[str, dict[str, str]]:
    """Mapea cada columna a su tipo, al estilo schemaMapping de RagaAI Catalyst."""
    return {
        column: {
            "columnType": infer_column_type(row.get(column) for row in rows),
        }
        for column in columns
    }


def _collect_columns(rows: list[dict[str, Any]], headers: Iterable[str] | None = None) -> list[str]:
    ordered: list[str] = []
    seen: set[str] = set()
    for name in _unique_names(headers or []):
        if name not in seen:
            ordered.append(name)
            seen.add(name)
    for row in rows:
        for key in row.keys():
            name = str(key)
            if name not in seen:
                ordered.append(name)
                seen.add(name)
    return ordered


def _tabular_payload(
    rows: list[dict[str, Any]],
    *,
    suffix: str,
    columns: list[str] | None = None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    serialized = [
        {str(key): _cell_value(value) for key, value in row.items()} for row in rows
    ]
    all_columns = _collect_columns(serialized, columns)
    return serialized, {
        "kind": "tabular",
        "format": suffix.removeprefix("."),
        "columns": all_columns,
        "schema_mapping": generate_schema(all_columns, serialized),
        "suggested_mapping": infer_mapping(all_columns),
    }


def _parse_csv(decoded: str) -> tuple[list[dict[str, Any]], list[str]]:
    sample = decoded[:8192]
    dialect = csv.excel
    if sample.strip():
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
    reader = csv.DictReader(io.StringIO(decoded), dialect=dialect)
    columns = _unique_names(reader.fieldnames or [])
    if reader.fieldnames:
        reader.fieldnames = columns
    rows: list[dict[str, Any]] = []
    for raw in reader:
        row: dict[str, Any] = {}
        for name in columns:
            row[name] = raw.get(name)
        for key, value in raw.items():
            if key in (None, ""):
                continue
            name = str(key)
            if name not in row:
                row[name] = value
                if name not in columns:
                    columns.append(name)
        rows.append(row)
    return rows, columns


def _json_to_rows(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        if not payload:
            return []
        if all(isinstance(item, dict) for item in payload):
            return payload
        return [{"value": item} for item in payload]
    if isinstance(payload, dict):
        for key in ("rows", "data", "items", "records"):
            nested = payload.get(key)
            if isinstance(nested, list):
                return _json_to_rows(nested)
        if payload and all(isinstance(value, list) for value in payload.values()):
            lengths = {len(value) for value in payload.values()}
            if len(lengths) == 1:
                size = lengths.pop()
                keys = list(payload.keys())
                return [
                    {key: payload[key][index] for key in keys} for index in range(size)
                ]
        return [payload]
    return [{"value": payload}]


def _parse_excel(content: bytes, suffix: str) -> tuple[list[dict[str, Any]], list[str]]:
    if suffix == ".xls":
        raise ValueError("Formato .xls no soportado; exporta a .xlsx o CSV")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ValueError("openpyxl es necesario para leer Excel") from exc
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        iterator = sheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if header_row is None:
            raise ValueError("Excel vacío")
        columns = _unique_names(header_row)
        rows: list[dict[str, Any]] = []
        for raw in iterator:
            if raw is None or all(cell in (None, "") for cell in raw):
                continue
            row: dict[str, Any] = {}
            for index, name in enumerate(columns):
                row[name] = raw[index] if index < len(raw) else None
            if len(raw) > len(columns):
                for index in range(len(columns), len(raw)):
                    extra = f"col_{index + 1}"
                    row[extra] = raw[index]
                    if extra not in columns:
                        columns.append(extra)
            rows.append(row)
    finally:
        workbook.close()
    return rows, columns


def _video_segments(content: bytes, filename: str) -> list[str]:
    """Transcribe el vídeo a fragmentos de habla. No modifica los parsers documentales."""
    return transcribe_video_segments(content, filename)


def _split_fragments(text: str, size: int = 1000, overlap: int = 200) -> list[str]:
    cleaned = re.sub(r"\r\n", "\n", text or "").strip()
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    if not cleaned:
        return []
    if len(cleaned) <= size:
        return [cleaned]
    parts: list[str] = []
    start = 0
    while start < len(cleaned):
        end = min(len(cleaned), start + size)
        if end < len(cleaned):
            window = cleaned[start:end]
            cut = window.rfind("\n\n")
            if cut < size * 0.4:
                cut = window.rfind("\n")
            if cut < size * 0.4:
                cut = window.rfind(" ")
            if cut >= size * 0.4:
                end = start + cut
        piece = cleaned[start:end].strip()
        if piece:
            parts.append(piece)
        if end >= len(cleaned):
            break
        start = max(end - overlap, start + 1)
    return parts


def _extract_office_xml_text(
    content: bytes,
    *,
    member: str | None = None,
    prefix: str | None = None,
) -> str:
    """Lee DOCX/PPTX desde el XML interno; no toca LlamaParse ni el factory."""
    try:
        from zipfile import ZipFile
        import xml.etree.ElementTree as ET

        with ZipFile(io.BytesIO(content)) as archive:
            names = [member] if member else sorted(
                name
                for name in archive.namelist()
                if prefix and name.startswith(prefix) and name.endswith(".xml")
            )
            parts: list[str] = []
            for name in names:
                if not name:
                    continue
                try:
                    root = ET.fromstring(archive.read(name))
                except KeyError:
                    continue
                texts = [
                    node.text
                    for node in root.iter()
                    if node.text and node.tag.endswith("}t")
                ]
                joined = "\n".join(texts).strip()
                if joined:
                    parts.append(joined)
            return "\n\n".join(parts).strip()
    except Exception:
        return ""


def _extract_document_text(
    filename: str, content: bytes, mime_type: str | None
) -> str:
    name = (filename or "").lower()
    mime = (mime_type or "").lower()
    if name.endswith(".docx") or "wordprocessingml" in mime:
        extracted = _extract_office_xml_text(content, member="word/document.xml")
        if extracted:
            return extracted
    if name.endswith((".pptx", ".ppt")) or "presentationml" in mime:
        extracted = _extract_office_xml_text(content, prefix="ppt/slides/slide")
        if extracted:
            return extracted
    from services.question_extraction import extract_text

    return extract_text(filename, content, mime_type).strip()


def _document_chunks(
    filename: str, content: bytes, mime_type: str | None
) -> tuple[list[dict], dict]:
    suffix = Path(filename or "").suffix.lower()
    is_video = is_video_file(filename, mime_type)
    try:
        if is_video:
            parts = _video_segments(content, filename)
        else:
            body = _extract_document_text(filename, content, mime_type).strip()
            parts = _split_fragments(body)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=422, detail=f"No se pudo leer el documento: {exc}"
        ) from exc
    if not parts:
        raise HTTPException(
            status_code=422,
            detail="No se pudo extraer texto ni fragmentos del archivo",
        )
    rows: list[dict[str, Any]] = []
    for index, part in enumerate(parts):
        headline = part.split("\n", 1)[0][:200]
        fragment = part.strip()
        rows.append(
            {
                "chunk_index": index,
                "headline": headline,
                "summary": fragment[:280],
                "fragment": fragment,
                "filename": filename,
                "metadata": {
                    "filename": filename,
                    "chunk_index": index,
                    "characters": len(fragment),
                    "source": "video" if is_video else suffix.removeprefix("."),
                },
            }
        )
    columns = [
        "chunk_index",
        "headline",
        "summary",
        "fragment",
        "filename",
        "metadata",
    ]
    return rows, {
        "kind": "chunks",
        "format": "video" if is_video else suffix.removeprefix(".") or "document",
        "columns": columns,
        "schema_mapping": generate_schema(columns, rows),
        "suggested_mapping": {
            "prompt": "headline",
            "context": "fragment",
            "expected_response": "summary",
            "metadata": "metadata",
        },
        "chunks": [
            {
                "chunk_index": row["chunk_index"],
                "headline": row["headline"],
                "summary": row["summary"],
                "fragment": row["fragment"],
                "filename": row["filename"],
                "title": row["filename"],
                "metadata": row["metadata"],
                "characters": len(str(row["fragment"])),
            }
            for row in rows
        ],
    }


def parse_content(
    filename: str, content: bytes, mime_type: str | None = None
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail="El archivo supera 100 MiB",
        )
    suffix = Path(filename or "").suffix.lower()
    if suffix in TABULAR_SUFFIXES:
        try:
            if suffix in {".xlsx", ".xls"}:
                rows, columns = _parse_excel(content, suffix)
            else:
                decoded = content.decode("utf-8-sig")
                if suffix == ".csv":
                    rows, columns = _parse_csv(decoded)
                elif suffix == ".json":
                    rows = _json_to_rows(json.loads(decoded))
                    columns = None
                else:
                    rows = [
                        json.loads(line) for line in decoded.splitlines() if line.strip()
                    ]
                    if not all(isinstance(item, dict) for item in rows):
                        raise ValueError("Cada línea JSONL debe ser un objeto")
                    columns = None
        except (UnicodeDecodeError, csv.Error, json.JSONDecodeError, ValueError) as exc:
            raise HTTPException(
                status_code=422, detail=f"Archivo inválido: {exc}"
            ) from exc
        if not rows and suffix != ".csv":
            raise HTTPException(status_code=422, detail="El archivo no contiene filas")
        payload, info = _tabular_payload(rows, suffix=suffix, columns=columns)
        if not info["columns"]:
            raise HTTPException(status_code=422, detail="No se detectaron columnas")
        return payload, info
    if suffix in LOG_SUFFIXES:
        try:
            rows = _parse_log_lines(content.decode("utf-8-sig"))
        except (UnicodeDecodeError, ValueError) as exc:
            raise HTTPException(
                status_code=422, detail=f"Archivo inválido: {exc}"
            ) from exc
        return _tabular_payload(rows, suffix=suffix)
    return _document_chunks(filename, content, mime_type)


async def preview_upload(file: UploadFile, limit: int = 20) -> dict[str, Any]:
    content = await file.read()
    rows, info = parse_content(file.filename or "upload", content, file.content_type)
    return {
        "filename": file.filename,
        "total_rows": len(rows),
        **info,
        "rows": rows[:limit],
    }


async def _validate_scope(
    db: AsyncSession,
    organization_id: str,
    tenant_id: str,
    knowledge_base_id: str | None,
) -> None:
    if not tenant_id:
        raise HTTPException(
            status_code=400, detail="La organización no tiene tenant activo"
        )
    valid_tenant = await db.scalar(
        text("""SELECT EXISTS (
                SELECT 1 FROM tenants
                WHERE id=:tenant AND organization_id=:org AND active=true
            )"""),
        {"tenant": tenant_id, "org": organization_id},
    )
    if not valid_tenant:
        raise HTTPException(status_code=403, detail="Tenant fuera de la organización")
    if knowledge_base_id:
        exists = await db.scalar(
            text("""SELECT EXISTS (
                    SELECT 1 FROM knowledge_bases
                    WHERE id=:kb AND tenant_id=:tenant
                )"""),
            {"kb": knowledge_base_id, "tenant": tenant_id},
        )
        if not exists:
            raise HTTPException(
                status_code=404, detail="Knowledge Base fuera del tenant"
            )


async def associate_departments(
    db: AsyncSession,
    *,
    organization_id: str,
    knowledge_base_id: str,
    department_ids: list[str],
    user_id: str,
    replace: bool = False,
) -> None:
    if replace:
        await db.execute(
            text(
                "DELETE FROM department_knowledge_bases "
                "WHERE knowledge_base_id=:kb AND organization_id=:org"
            ),
            {"kb": knowledge_base_id, "org": organization_id},
        )
    for department_id in dict.fromkeys(department_ids):
        valid = await db.scalar(
            text(
                "SELECT EXISTS (SELECT 1 FROM departments "
                "WHERE id=:department AND organization_id=:org)"
            ),
            {"department": department_id, "org": organization_id},
        )
        if not valid:
            raise HTTPException(
                status_code=404, detail=f"Departamento no válido: {department_id}"
            )
        await db.execute(
            text("""INSERT INTO department_knowledge_bases
                   (department_id, knowledge_base_id, organization_id, created_by)
                   VALUES (:department, :kb, :org, :user)
                   ON CONFLICT (department_id, knowledge_base_id) DO NOTHING"""),
            {
                "department": department_id,
                "kb": knowledge_base_id,
                "org": organization_id,
                "user": user_id,
            },
        )


async def associate_dataset_departments(
    db: AsyncSession,
    *,
    dataset_id: str,
    organization_id: str,
    department_ids: list[str],
) -> None:
    await db.execute(
        text("DELETE FROM rag_dataset_departments WHERE dataset_id=:dataset"),
        {"dataset": dataset_id},
    )
    for department_id in dict.fromkeys(department_ids):
        exists = await db.scalar(
            text(
                "SELECT EXISTS (SELECT 1 FROM departments "
                "WHERE id=:department AND organization_id=:org)"
            ),
            {"department": department_id, "org": organization_id},
        )
        if not exists:
            raise HTTPException(
                status_code=404, detail=f"Departamento no válido: {department_id}"
            )
        await db.execute(
            text("""INSERT INTO rag_dataset_departments (dataset_id, department_id)
                   VALUES (:dataset, :department)
                   ON CONFLICT DO NOTHING"""),
            {"dataset": dataset_id, "department": department_id},
        )


async def import_rows(
    db: AsyncSession,
    *,
    rows: list[dict[str, Any]],
    mapping: dict[str, str],
    name: str,
    description: str | None,
    source_type: str,
    source_name: str | None,
    source_metadata: dict[str, Any],
    organization_id: str,
    tenant_id: str,
    knowledge_base_id: str | None,
    department_ids: list[str],
    user_id: str,
) -> dict[str, Any]:
    await _validate_scope(db, organization_id, tenant_id, knowledge_base_id)
    normalized: list[dict[str, Any]] = []
    rejected = 0
    seen: set[str] = set()
    for row in rows:
        try:
            item = normalize_row(row, mapping)
        except ValueError:
            rejected += 1
            continue
        if item["code_hash"] in seen:
            rejected += 1
            continue
        seen.add(item["code_hash"])
        normalized.append(item)
    if not normalized:
        raise HTTPException(
            status_code=422, detail="No hay filas RAG válidas para importar"
        )

    dataset_id = str(uuid4())
    await db.execute(
        text("""INSERT INTO rag_datasets
               (id, organization_id, tenant_id, knowledge_base_id, name, description,
                source_type, source_name, source_metadata, mapping, status, row_count, created_by)
               VALUES (:id, :org, :tenant, :kb, :name, :description, :source_type,
                       :source_name, CAST(:source_metadata AS jsonb), CAST(:mapping AS jsonb),
                       'ready', :row_count, :user)"""),
        {
            "id": dataset_id,
            "org": organization_id,
            "tenant": tenant_id,
            "kb": knowledge_base_id,
            "name": name,
            "description": description,
            "source_type": source_type,
            "source_name": source_name,
            "source_metadata": json.dumps(source_metadata, default=str),
            "mapping": json.dumps(mapping),
            "row_count": len(normalized),
            "user": user_id,
        },
    )
    for item in normalized:
        await db.execute(
            text("""INSERT INTO rag_dataset_rows
                   (dataset_id, organization_id, tenant_id, knowledge_base_id, source,
                    status, "traceId", prompt, context, response, timestamp,
                    expected_response, expected_context, metadata, pipeline,
                    alternate_response, code_hash)
                   VALUES (:dataset, :org, :tenant, :kb, :source, :status, :trace_id,
                           :prompt, CAST(:context AS jsonb), :response, :timestamp,
                           :expected_response, CAST(:expected_context AS jsonb),
                           CAST(:metadata AS jsonb), CAST(:pipeline AS jsonb),
                           :alternate_response, :code_hash)
                   ON CONFLICT (dataset_id, code_hash) DO NOTHING"""),
            {
                "dataset": dataset_id,
                "org": organization_id,
                "tenant": tenant_id,
                "kb": knowledge_base_id,
                "source": source_type,
                "status": "pending_review" if source_type == "synthetic" else "ready",
                "trace_id": item["traceId"],
                "prompt": item["prompt"],
                "context": json.dumps(item["context"], default=str),
                "response": item["response"],
                "timestamp": item["timestamp"],
                "expected_response": item["expected_response"],
                "expected_context": json.dumps(item["expected_context"], default=str),
                "metadata": json.dumps(item["metadata"], default=str),
                "pipeline": json.dumps(item["pipeline"], default=str),
                "alternate_response": item["alternate_response"],
                "code_hash": item["code_hash"],
            },
        )
    if department_ids:
        await associate_dataset_departments(
            db,
            dataset_id=dataset_id,
            organization_id=organization_id,
            department_ids=department_ids,
        )
    if knowledge_base_id and department_ids:
        await associate_departments(
            db,
            organization_id=organization_id,
            knowledge_base_id=knowledge_base_id,
            department_ids=department_ids,
            user_id=user_id,
        )
    await db.commit()
    try:
        from services.rag_service import RAGService

        await RAGService().notify_index_updated(str(tenant_id))
    except Exception:
        pass
    return {
        "id": dataset_id,
        "name": name,
        "status": "ready",
        "source_type": source_type,
        "imported_rows": len(normalized),
        "rejected_rows": rejected,
        "department_ids": department_ids,
    }


async def synthetic_rows(
    db: AsyncSession,
    *,
    organization_id: str,
    tenant_id: str,
    knowledge_base_id: str | None,
    document_ids: list[str],
    questions: list[str],
    limit: int,
) -> list[dict[str, Any]]:
    clauses = ["d.tenant_id = :tenant"]
    params: dict[str, Any] = {"tenant": tenant_id, "limit": limit}
    if knowledge_base_id:
        clauses.append("dq.knowledge_base_id = :kb")
        params["kb"] = knowledge_base_id
    if document_ids:
        clauses.append("dq.document_id = ANY(CAST(:documents AS uuid[]))")
        params["documents"] = document_ids
    result = await db.execute(
        text(f"""SELECT dq.question, dq.reference_answer, dq.rationale, dq.category,
                       dq.document_id, dq.filename,
                       COALESCE((
                         SELECT jsonb_agg(c.content ORDER BY c.position)
                         FROM chunks c WHERE c.document_id=dq.document_id
                       ), '[]'::jsonb) AS contexts
                FROM document_questions dq
                JOIN documents d ON d.id=dq.document_id
                WHERE {' AND '.join(clauses)}
                ORDER BY dq.created_at DESC LIMIT :limit"""),
        params,
    )
    rows: list[dict[str, Any]] = []
    seen_prompts: set[str] = set()
    for row in result.mappings().all():
        prompt = str(row["question"]).strip()
        if not prompt or prompt.casefold() in seen_prompts:
            continue
        seen_prompts.add(prompt.casefold())
        rows.append(
            {
                "prompt": prompt,
                "context": list(row.get("contexts") or [])[:8],
                "expected_response": row.get("reference_answer")
                or row.get("rationale"),
                "metadata": {
                    "document_id": str(row["document_id"]),
                    "filename": row.get("filename"),
                    "category": row.get("category"),
                },
            }
        )
    if len(rows) < limit:
        document_clauses = ["d.tenant_id=:tenant"]
        document_params: dict[str, Any] = {
            "tenant": tenant_id,
            "remaining": limit - len(rows),
        }
        if knowledge_base_id:
            document_clauses.append("d.knowledge_base_id=:kb")
            document_params["kb"] = knowledge_base_id
        if document_ids:
            document_clauses.append("d.id = ANY(CAST(:documents AS uuid[]))")
            document_params["documents"] = document_ids
        documents = await db.execute(
            text(f"""SELECT d.id, d.filename, d.title, d.description,
                           COALESCE((
                             SELECT jsonb_agg(c.content ORDER BY c.position)
                             FROM chunks c WHERE c.document_id=d.id
                           ), '[]'::jsonb) AS contexts
                    FROM documents d
                    WHERE {' AND '.join(document_clauses)}
                    ORDER BY d.uploaded_at DESC LIMIT :remaining"""),
            document_params,
        )
        for document in documents.mappings().all():
            contexts = list(document.get("contexts") or [])[:8]
            if not contexts:
                continue
            title = document.get("title") or document.get("filename") or "el documento"
            prompt = f"¿Qué información relevante contiene {title}?"
            if prompt.casefold() in seen_prompts:
                continue
            seen_prompts.add(prompt.casefold())
            rows.append(
                {
                    "prompt": prompt,
                    "context": contexts,
                    "expected_response": document.get("description")
                    or contexts[0][:1000],
                    "metadata": {
                        "document_id": str(document["id"]),
                        "filename": document.get("filename"),
                        "origin": "document",
                    },
                }
            )
    for question in questions:
        prompt = question.strip()
        if prompt and prompt.casefold() not in seen_prompts:
            seen_prompts.add(prompt.casefold())
            rows.append({"prompt": prompt, "metadata": {"origin": "manual"}})
    if not rows:
        topic = next((q.strip() for q in questions if q.strip()), "la explotación")
        templates = [
            f"¿Qué debo revisar primero ante un problema de {topic}?",
            f"¿Qué herramientas o recambios uso para {topic}?",
            f"¿Qué precauciones de seguridad aplican a {topic}?",
            f"¿Qué datos técnicos (dosis, tiempos, temperaturas) hay que tener en cuenta en {topic}?",
            f"¿Cuándo debo llamar a un técnico si falla {topic}?",
        ]
        for prompt in templates[:limit]:
            rows.append(
                {
                    "prompt": prompt,
                    "expected_response": (
                        "Respuesta de orientación: confirma el síntoma en campo, "
                        "consulta la ficha técnica y no improvises materias activas."
                    ),
                    "metadata": {"origin": "synthetic_template", "topic": topic},
                }
            )
    return rows[:limit]


async def enqueue_synthetic_hitl(
    db: AsyncSession,
    *,
    rows: list[dict[str, Any]],
    organization_id: str,
    knowledge_base_id: str | None,
    user_id: str,
) -> int:
    queued = 0
    for row in rows:
        prompt = str(row.get("prompt") or "").strip()
        if not prompt:
            continue
        exists = await db.scalar(
            text("""SELECT EXISTS (
                    SELECT 1 FROM rag_human_reviews
                    WHERE organization_id=:org AND source='synthetic_dataset'
                      AND lower(question)=lower(:question)
                )"""),
            {"org": organization_id, "question": prompt},
        )
        if exists:
            continue
        await db.execute(
            text("""INSERT INTO rag_human_reviews
                   (id, source, user_id, organization_id, knowledge_base_id,
                    document_id, question, answer, context_snippet, status)
                   VALUES (:id, 'synthetic_dataset', :user, :org, :kb, :document,
                           :question, :answer, :context, 'pending')"""),
            {
                "id": str(uuid4()),
                "user": user_id,
                "org": organization_id,
                "kb": knowledge_base_id,
                "document": (row.get("metadata") or {}).get("document_id"),
                "question": prompt,
                "answer": row.get("expected_response"),
                "context": "\n---\n".join(row.get("context") or [])[:4000] or None,
            },
        )
        queued += 1
    await db.commit()
    return queued


async def _dataset_in_scope(
    db: AsyncSession,
    *,
    dataset_id: str,
    organization_id: str,
    tenant_id: str,
) -> dict[str, Any]:
    result = await db.execute(
        text(
            """SELECT id, name, description, knowledge_base_id, mapping, status,
                      organization_id, tenant_id, row_count, source_type
               FROM rag_datasets
               WHERE id=:id AND organization_id=:org AND tenant_id=:tenant"""
        ),
        {"id": dataset_id, "org": organization_id, "tenant": tenant_id},
    )
    dataset = result.mappings().first()
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset no encontrado")
    return dict(dataset)


async def _notify_index(tenant_id: str) -> None:
    try:
        from services.rag_service import RAGService

        await RAGService().notify_index_updated(str(tenant_id))
    except Exception:
        pass


async def update_dataset(
    db: AsyncSession,
    *,
    dataset_id: str,
    organization_id: str,
    tenant_id: str,
    user_id: str,
    name: str | None = None,
    description: str | None = None,
    knowledge_base_id: str | None = None,
    department_ids: list[str] | None = None,
    mapping: dict[str, str] | None = None,
) -> dict[str, Any]:
    current = await _dataset_in_scope(
        db,
        dataset_id=dataset_id,
        organization_id=organization_id,
        tenant_id=tenant_id,
    )
    next_name = (name or "").strip() or current["name"]
    next_description = current["description"] if description is None else description
    next_kb = (
        current["knowledge_base_id"]
        if knowledge_base_id is None
        else (knowledge_base_id or None)
    )
    if next_kb:
        next_kb = str(next_kb)
        await _validate_scope(db, organization_id, tenant_id, next_kb)
    next_mapping = current["mapping"] if mapping is None else mapping
    if mapping is not None and not isinstance(mapping, dict):
        raise HTTPException(status_code=422, detail="mapping debe ser un objeto")
    await db.execute(
        text(
            """UPDATE rag_datasets
               SET name=:name,
                   description=:description,
                   knowledge_base_id=:kb,
                   mapping=CAST(:mapping AS jsonb),
                   updated_at=CURRENT_TIMESTAMP
               WHERE id=:id AND organization_id=:org AND tenant_id=:tenant"""
        ),
        {
            "name": next_name,
            "description": next_description,
            "kb": next_kb,
            "mapping": json.dumps(next_mapping or {}, default=str),
            "id": dataset_id,
            "org": organization_id,
            "tenant": tenant_id,
        },
    )
    if knowledge_base_id is not None:
        await db.execute(
            text(
                """UPDATE rag_dataset_rows
                   SET knowledge_base_id=:kb
                   WHERE dataset_id=:id AND organization_id=:org AND tenant_id=:tenant"""
            ),
            {
                "kb": next_kb,
                "id": dataset_id,
                "org": organization_id,
                "tenant": tenant_id,
            },
        )
    if department_ids is not None:
        await associate_dataset_departments(
            db,
            dataset_id=dataset_id,
            organization_id=organization_id,
            department_ids=department_ids,
        )
        if next_kb and department_ids:
            await associate_departments(
                db,
                organization_id=organization_id,
                knowledge_base_id=str(next_kb),
                department_ids=department_ids,
                user_id=user_id,
            )
    await db.commit()
    await _notify_index(tenant_id)
    return {
        "id": dataset_id,
        "name": next_name,
        "description": next_description,
        "knowledge_base_id": next_kb,
        "department_ids": department_ids,
        "status": current["status"],
    }


async def delete_dataset(
    db: AsyncSession,
    *,
    dataset_id: str,
    organization_id: str,
    tenant_id: str,
) -> dict[str, Any]:
    await _dataset_in_scope(
        db,
        dataset_id=dataset_id,
        organization_id=organization_id,
        tenant_id=tenant_id,
    )
    await db.execute(
        text(
            """DELETE FROM rag_datasets
               WHERE id=:id AND organization_id=:org AND tenant_id=:tenant"""
        ),
        {"id": dataset_id, "org": organization_id, "tenant": tenant_id},
    )
    await db.commit()
    await _notify_index(tenant_id)
    return {"id": dataset_id, "deleted": True}


def _as_mapping(value: Any) -> dict[str, str]:
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except json.JSONDecodeError:
            return {}
    if not isinstance(value, dict):
        return {}
    return {str(key): str(item) for key, item in value.items() if item}


async def append_rows(
    db: AsyncSession,
    *,
    dataset_id: str,
    rows: list[dict[str, Any]],
    organization_id: str,
    tenant_id: str,
    source_name: str | None,
    mapping: dict[str, str] | None = None,
) -> dict[str, Any]:
    """Añade filas a un dataset ya ingestado (equivalente a Dataset.add_rows de Catalyst)."""
    current = await _dataset_in_scope(
        db,
        dataset_id=dataset_id,
        organization_id=organization_id,
        tenant_id=tenant_id,
    )
    mapping = mapping or _as_mapping(current.get("mapping"))
    if not mapping.get("prompt"):
        mapping = {**infer_mapping(rows[0].keys() if rows else []), **mapping}
    knowledge_base_id = current.get("knowledge_base_id")
    if knowledge_base_id:
        knowledge_base_id = str(knowledge_base_id)
    normalized: list[dict[str, Any]] = []
    rejected = 0
    seen: set[str] = set()
    for row in rows:
        try:
            item = normalize_row(row, mapping)
        except ValueError:
            rejected += 1
            continue
        if item["code_hash"] in seen:
            rejected += 1
            continue
        seen.add(item["code_hash"])
        normalized.append(item)
    if not normalized:
        raise HTTPException(
            status_code=422, detail="No hay filas RAG válidas para añadir"
        )
    inserted = 0
    for item in normalized:
        result = await db.execute(
            text("""INSERT INTO rag_dataset_rows
                   (dataset_id, organization_id, tenant_id, knowledge_base_id, source,
                    status, "traceId", prompt, context, response, timestamp,
                    expected_response, expected_context, metadata, pipeline,
                    alternate_response, code_hash)
                   VALUES (:dataset, :org, :tenant, :kb, :source, :status, :trace_id,
                           :prompt, CAST(:context AS jsonb), :response, :timestamp,
                           :expected_response, CAST(:expected_context AS jsonb),
                           CAST(:metadata AS jsonb), CAST(:pipeline AS jsonb),
                           :alternate_response, :code_hash)
                   ON CONFLICT (dataset_id, code_hash) DO NOTHING"""),
            {
                "dataset": dataset_id,
                "org": organization_id,
                "tenant": tenant_id,
                "kb": knowledge_base_id,
                "source": current.get("source_type") or "upload",
                "status": "ready",
                "trace_id": item["traceId"],
                "prompt": item["prompt"],
                "context": json.dumps(item["context"], default=str),
                "response": item["response"],
                "timestamp": item["timestamp"],
                "expected_response": item["expected_response"],
                "expected_context": json.dumps(item["expected_context"], default=str),
                "metadata": json.dumps(item["metadata"], default=str),
                "pipeline": json.dumps(item["pipeline"], default=str),
                "alternate_response": item["alternate_response"],
                "code_hash": item["code_hash"],
            },
        )
        inserted += int(result.rowcount or 0)
    count = await db.scalar(
        text("SELECT COUNT(*) FROM rag_dataset_rows WHERE dataset_id=:id"),
        {"id": dataset_id},
    )
    await db.execute(
        text(
            """UPDATE rag_datasets
               SET row_count=:count, updated_at=CURRENT_TIMESTAMP
               WHERE id=:id"""
        ),
        {"count": int(count or 0), "id": dataset_id},
    )
    await db.commit()
    await _notify_index(tenant_id)
    return {
        "id": dataset_id,
        "name": current["name"],
        "imported_rows": inserted,
        "rejected_rows": rejected + (len(normalized) - inserted),
        "row_count": int(count or 0),
        "source_name": source_name,
    }


_PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}")
_PII_PATTERNS = (
    (re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I), "email"),
    (re.compile(r"\b\d{8}[A-Z]\b"), "nif"),
    (re.compile(r"\b(?:\+34)?[6-9]\d{8}\b"), "phone"),
)
DEFAULT_LLM_COLUMN_TEMPLATE = (
    "Evalúa en una frase si la respuesta está fundada en el contexto.\n"
    "Pregunta: {{prompt}}\n"
    "Contexto: {{context}}\n"
    "Respuesta: {{response}}"
)
_TRACE_MAPPING = {
    "prompt": "prompt",
    "response": "response",
    "traceId": "traceId",
    "timestamp": "timestamp",
    "metadata": "metadata",
}


def _as_dict(value: Any) -> dict[str, Any]:
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except json.JSONDecodeError:
            return {}
    return dict(value) if isinstance(value, dict) else {}


def context_list(value: Any) -> list[str]:
    parsed = value
    if isinstance(parsed, str):
        try:
            parsed = json.loads(parsed)
        except json.JSONDecodeError:
            return [parsed] if parsed.strip() else []
    if isinstance(parsed, list):
        return [str(item) for item in parsed if item not in (None, "")]
    if parsed in (None, ""):
        return []
    if isinstance(parsed, dict):
        return [json.dumps(parsed, ensure_ascii=False)]
    return [str(parsed)]


def row_text_fields(row: dict[str, Any]) -> dict[str, str]:
    return {
        "prompt": str(row.get("prompt") or ""),
        "context": "\n".join(context_list(row.get("context"))),
        "response": str(row.get("response") or ""),
        "expected_response": str(row.get("expected_response") or ""),
        "traceId": str(row.get("traceId") or ""),
    }


def render_template(template: str, fields: dict[str, str]) -> str:
    return _PLACEHOLDER.sub(lambda match: fields.get(match.group(1), ""), template)


def sanitize_column_name(value: str) -> str:
    name = re.sub(r"\s+", "_", (value or "").strip())
    name = re.sub(r"[^A-Za-z0-9_-]", "", name)
    if not name:
        raise HTTPException(status_code=422, detail="Nombre de columna inválido")
    return name[:80]


def pair_chat_messages(
    messages: list[dict[str, Any]], limit: int = 50
) -> list[dict[str, Any]]:
    pairs: list[dict[str, Any]] = []
    index = 0
    while index < len(messages) - 1 and len(pairs) < limit:
        current = messages[index]
        nxt = messages[index + 1]
        same_conversation = str(current.get("conversation_id")) == str(
            nxt.get("conversation_id")
        )
        if (
            str(current.get("role") or "").lower() == "user"
            and str(nxt.get("role") or "").lower() == "assistant"
            and same_conversation
        ):
            created = nxt.get("created_at")
            timestamp = (
                created.isoformat() if hasattr(created, "isoformat") else created
            )
            pairs.append(
                {
                    "traceId": str(nxt.get("id") or ""),
                    "prompt": str(current.get("content") or ""),
                    "response": str(nxt.get("content") or ""),
                    "timestamp": timestamp,
                    "metadata": {
                        "origin": "chat_trace",
                        "conversation_id": str(current.get("conversation_id") or ""),
                        "prompt_tokens": nxt.get("prompt_tokens"),
                        "completion_tokens": nxt.get("completion_tokens"),
                    },
                }
            )
            index += 2
            continue
        index += 1
    return pairs


def traces_from_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    traces: list[dict[str, Any]] = []
    for row in rows:
        metadata = _as_dict(row.get("metadata"))
        traces.append(
            {
                "traceId": str(row.get("traceId") or row.get("id") or ""),
                "row_id": str(row.get("id") or ""),
                "timestamp": row.get("timestamp") or row.get("created_at"),
                "prompt": row.get("prompt"),
                "response": row.get("response"),
                "context": row.get("context"),
                "spans": [
                    {"name": "user", "content": str(row.get("prompt") or "")},
                    {
                        "name": "retrieve",
                        "content": "\n".join(context_list(row.get("context"))),
                    },
                    {"name": "generate", "content": str(row.get("response") or "")},
                ],
                "guardrails": metadata.get("guardrails") or {},
                "llm_columns": metadata.get("llm_columns") or {},
            }
        )
    return traces


def evaluate_row_guardrails(row: dict[str, Any]) -> dict[str, Any]:
    from evaluation.hallucination import decide_is_hallucination

    fields = row_text_fields(row)
    contexts = context_list(row.get("context"))
    flags: list[str] = []
    details: dict[str, Any] = {}
    if not fields["context"].strip():
        flags.append("missing_context")
    if not fields["response"].strip():
        flags.append("missing_response")
    hallucination = decide_is_hallucination(
        answer=fields["response"],
        contexts=contexts,
    )
    if hallucination.get("is_hallucination"):
        flags.append("hallucination")
        details["hallucination"] = hallucination
    haystack = " ".join(
        (fields["prompt"], fields["context"], fields["response"])
    )
    pii_hits = [
        kind for pattern, kind in _PII_PATTERNS if pattern.search(haystack)
    ]
    if pii_hits:
        flags.append("pii")
        details["pii"] = pii_hits
    return {
        "passed": not flags,
        "flags": flags,
        "details": details,
    }


async def llm_complete(prompt: str) -> str:
    from rag.adapters.outbound.ollama import OllamaLlmAdapter

    model = os.getenv("RAG_GENERATION_MODEL", "llama3.2:latest")
    adapter = OllamaLlmAdapter()
    return (
        await adapter.complete(
            model,
            [
                {
                    "role": "system",
                    "content": "Responde de forma breve y factual en español.",
                },
                {"role": "user", "content": prompt},
            ],
        )
    ).strip()


async def _fetch_dataset_rows(
    db: AsyncSession,
    *,
    dataset_id: str,
    organization_id: str,
    tenant_id: str,
    limit: int,
) -> list[dict[str, Any]]:
    result = await db.execute(
        text(
            """SELECT id, "traceId", prompt, context, response, timestamp,
                      expected_response, expected_context, metadata, pipeline,
                      alternate_response, code_hash, source, status, created_at
               FROM rag_dataset_rows
               WHERE dataset_id=:id AND organization_id=:org AND tenant_id=:tenant
               ORDER BY created_at, id LIMIT :limit"""
        ),
        {
            "id": dataset_id,
            "org": organization_id,
            "tenant": tenant_id,
            "limit": limit,
        },
    )
    return [dict(row) for row in result.mappings().all()]


async def _update_row_metadata(
    db: AsyncSession,
    *,
    row_id: Any,
    metadata: dict[str, Any],
) -> None:
    await db.execute(
        text(
            """UPDATE rag_dataset_rows
               SET metadata=CAST(:metadata AS jsonb)
               WHERE id=:id"""
        ),
        {"id": row_id, "metadata": json.dumps(metadata, default=str)},
    )


async def list_traces(
    db: AsyncSession,
    *,
    dataset_id: str,
    organization_id: str,
    tenant_id: str,
    limit: int = 100,
) -> dict[str, Any]:
    await _dataset_in_scope(
        db,
        dataset_id=dataset_id,
        organization_id=organization_id,
        tenant_id=tenant_id,
    )
    rows = await _fetch_dataset_rows(
        db,
        dataset_id=dataset_id,
        organization_id=organization_id,
        tenant_id=tenant_id,
        limit=limit,
    )
    traces = traces_from_rows(rows)
    return {"dataset_id": dataset_id, "count": len(traces), "data": traces}


async def import_traces_from_chat(
    db: AsyncSession,
    *,
    dataset_id: str,
    organization_id: str,
    tenant_id: str,
    limit: int = 50,
) -> dict[str, Any]:
    current = await _dataset_in_scope(
        db,
        dataset_id=dataset_id,
        organization_id=organization_id,
        tenant_id=tenant_id,
    )
    knowledge_base_id = current.get("knowledge_base_id")
    result = await db.execute(
        text(
            """SELECT m.id, m.role, m.content, m.created_at, m.conversation_id,
                      m.prompt_tokens, m.completion_tokens
               FROM messages m
               JOIN conversations c ON c.id = m.conversation_id
               WHERE c.tenant_id=:tenant
                 AND (:kb IS NULL OR c.knowledge_base_id=:kb)
               ORDER BY m.conversation_id, m.created_at
               LIMIT :limit"""
        ),
        {
            "tenant": tenant_id,
            "kb": str(knowledge_base_id) if knowledge_base_id else None,
            "limit": max(limit * 2, 20),
        },
    )
    pairs = pair_chat_messages(
        [dict(row) for row in result.mappings().all()],
        limit=limit,
    )
    if not pairs:
        raise HTTPException(
            status_code=422,
            detail="No hay traces de chat para importar en este ámbito",
        )
    return await append_rows(
        db,
        dataset_id=dataset_id,
        rows=pairs,
        organization_id=organization_id,
        tenant_id=tenant_id,
        source_name="chat_traces",
        mapping=_TRACE_MAPPING,
    )


async def add_llm_column(
    db: AsyncSession,
    *,
    dataset_id: str,
    organization_id: str,
    tenant_id: str,
    column_name: str,
    prompt_template: str,
    limit: int = 25,
    completer: Callable[[str], Awaitable[str]] | None = None,
) -> dict[str, Any]:
    name = sanitize_column_name(column_name)
    template = (prompt_template or "").strip() or DEFAULT_LLM_COLUMN_TEMPLATE
    await _dataset_in_scope(
        db,
        dataset_id=dataset_id,
        organization_id=organization_id,
        tenant_id=tenant_id,
    )
    rows = await _fetch_dataset_rows(
        db,
        dataset_id=dataset_id,
        organization_id=organization_id,
        tenant_id=tenant_id,
        limit=limit,
    )
    if not rows:
        raise HTTPException(status_code=422, detail="El dataset no tiene filas")
    complete = completer or llm_complete
    filled = 0
    errors = 0
    for row in rows:
        prompt = render_template(template, row_text_fields(row))
        try:
            value = await complete(prompt)
        except Exception as exc:
            value = f"(error: {exc})"
            errors += 1
        metadata = _as_dict(row.get("metadata"))
        columns = dict(metadata.get("llm_columns") or {})
        columns[name] = value
        metadata["llm_columns"] = columns
        await _update_row_metadata(db, row_id=row["id"], metadata=metadata)
        filled += 1
    await db.commit()
    return {
        "id": dataset_id,
        "column_name": name,
        "filled_rows": filled,
        "errors": errors,
    }


async def run_guardrails(
    db: AsyncSession,
    *,
    dataset_id: str,
    organization_id: str,
    tenant_id: str,
    limit: int = 100,
) -> dict[str, Any]:
    await _dataset_in_scope(
        db,
        dataset_id=dataset_id,
        organization_id=organization_id,
        tenant_id=tenant_id,
    )
    rows = await _fetch_dataset_rows(
        db,
        dataset_id=dataset_id,
        organization_id=organization_id,
        tenant_id=tenant_id,
        limit=limit,
    )
    if not rows:
        raise HTTPException(status_code=422, detail="El dataset no tiene filas")
    flag_counts: dict[str, int] = {}
    passed = 0
    results: list[dict[str, Any]] = []
    for row in rows:
        verdict = evaluate_row_guardrails(row)
        metadata = _as_dict(row.get("metadata"))
        metadata["guardrails"] = verdict
        await _update_row_metadata(db, row_id=row["id"], metadata=metadata)
        if verdict["passed"]:
            passed += 1
        for flag in verdict["flags"]:
            flag_counts[flag] = flag_counts.get(flag, 0) + 1
        results.append(
            {
                "row_id": str(row.get("id") or ""),
                "traceId": str(row.get("traceId") or ""),
                **verdict,
            }
        )
    await db.commit()
    scanned = len(rows)
    return {
        "id": dataset_id,
        "scanned": scanned,
        "passed": passed,
        "flagged": scanned - passed,
        "flags": flag_counts,
        "rows": results,
    }
